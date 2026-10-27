"""DOCX: adaptado de RAG STUDIO (infrastructure/extraction/
text_extractors.py: `DocxExtractor`, `_iter_block_items`,
`_table_to_markdown`) -- copiado y ajustado como codigo propio (ADR-001,
seccion 8), no importado en vivo. Los parrafos con estilo 'Heading N' o
'Title' se normalizan a encabezados Markdown; las tablas se convierten
a tablas Markdown en el mismo orden en que aparecen en el documento.
"""
from __future__ import annotations

from collections.abc import Callable

import docx
import docx.document
from docx.oxml.ns import qn
from docx.table import Table as DocxTable
from docx.text.paragraph import Paragraph as DocxParagraph

from domain.documents.entities import FormatoDocumento, ResultadoExtraccion, TipoProcesamiento
from domain.documents.exceptions import ExtractionFailedError
from domain.documents.ports import DocumentExtractor

from .markdown_helpers import rows_to_markdown_table


def _iter_block_items(document: "docx.document.Document"):
    """Recorre parrafos y tablas de un DOCX en su orden real de
    aparicion -- python-docx no lo expone directamente
    (`document.paragraphs`/`document.tables` son dos listas separadas,
    sin orden relativo entre si); idiom estandar de la comunidad
    python-docx."""
    for child in document.element.body.iterchildren():
        if child.tag == qn("w:p"):
            yield DocxParagraph(child, document)
        elif child.tag == qn("w:tbl"):
            yield DocxTable(child, document)


def _heading_level(style_name: str) -> int:
    style_name = (style_name or "").strip()
    if style_name == "Title":
        return 1
    if style_name.startswith("Heading "):
        suffix = style_name[len("Heading "):].strip()
        if suffix.isdigit():
            return min(int(suffix), 6)
    return 0


def _table_to_markdown(table: DocxTable) -> str:
    rows = [[cell.text for cell in row.cells] for row in table.rows]
    return rows_to_markdown_table(rows)


class DocxExtractor(DocumentExtractor):
    def supports(self, formato: FormatoDocumento, tipo_procesamiento: TipoProcesamiento) -> bool:
        # DOCX siempre es EXTRACCION_DIRECTA -- el parametro se ignora.
        return formato == FormatoDocumento.DOCX

    def extract(
        self,
        ruta: str,
        formato: FormatoDocumento,
        tipo_procesamiento: TipoProcesamiento,
        on_progreso: Callable[[int, int], None] | None = None,
    ) -> ResultadoExtraccion:
        # on_progreso se ignora: extraccion directa de DOCX es rapida
        # incluso en documentos grandes (sin OCR de por medio).
        try:
            document = docx.Document(ruta)
        except Exception as exc:  # python-docx no tiene una excepcion propia especifica para "archivo invalido"
            raise ExtractionFailedError(ruta, str(exc)) from exc

        lines: list[str] = []
        for block in _iter_block_items(document):
            if isinstance(block, DocxTable):
                tabla = _table_to_markdown(block)
                if tabla:
                    lines.append(tabla)
                continue
            texto = block.text.strip()
            if not texto:
                continue
            nivel = _heading_level(block.style.name if block.style else "")
            lines.append(f"{'#' * nivel} {texto}" if nivel else texto)

        advertencias = [] if lines else ["El documento DOCX no tiene parrafos ni tablas con texto."]
        return ResultadoExtraccion(markdown="\n\n".join(lines), metodo="docx", advertencias=advertencias)
