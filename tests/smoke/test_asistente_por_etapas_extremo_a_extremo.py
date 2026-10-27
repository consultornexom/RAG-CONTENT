"""Prueba de humo de la Fase 11.2 (Asistente guiado por etapas,
Documento Maestro seccion 4.11.1): recorre el asistente de punta a
punta -- bootstrap()/abrir_proyecto() reales, sin fakes, mismo patron
que tests/smoke/test_gui_menu_completo_extremo_a_extremo.py (Fase 11).
Los dialogos modales (CrearProyectoDialog) y los QMessageBox de
resultado se parchan igual que alli, para no bloquear la prueba en modo
headless."""
from __future__ import annotations

import shutil
from pathlib import Path

import docx
import pytest

pytest.importorskip("PySide6")

from PySide6.QtCore import Qt  # noqa: E402
from PySide6.QtWidgets import QFileDialog, QMessageBox  # noqa: E402

from app.bootstrap import bootstrap  # noqa: E402
from domain.documents.entities import EstadoDocumento  # noqa: E402
from ui.dialogs import CrearProyectoDialog  # noqa: E402
from ui.main_window import MainWindow  # noqa: E402

RAIZ_PRUEBA = Path("/tmp/rag_content_studio_smoke_test_asistente")


@pytest.fixture()
def raiz_temporal():
    if RAIZ_PRUEBA.exists():
        shutil.rmtree(RAIZ_PRUEBA)
    RAIZ_PRUEBA.mkdir(parents=True)
    perfiles_dir = RAIZ_PRUEBA / "perfiles" / "perfil_generico"
    perfiles_dir.mkdir(parents=True)
    shutil.copy(
        Path(__file__).resolve().parents[2] / "perfiles" / "perfil_generico" / "0.1.yaml",
        perfiles_dir / "0.1.yaml",
    )
    origenes_dir = RAIZ_PRUEBA / "_origenes_de_prueba"
    origenes_dir.mkdir()
    yield RAIZ_PRUEBA, origenes_dir
    shutil.rmtree(RAIZ_PRUEBA, ignore_errors=True)


def _crear_proyecto_desde_asistente(ventana: MainWindow, monkeypatch, nombre: str) -> None:
    def _prellenado(self, parent=None, perfiles=None):
        CrearProyectoDialog.__base_init__(self, parent, perfiles)
        self.campo_nombre.setText(nombre)
        indice = self.campo_perfil.findData("perfil_generico")
        self.campo_perfil.setCurrentIndex(indice)

    CrearProyectoDialog.__base_init__ = CrearProyectoDialog.__init__
    monkeypatch.setattr(CrearProyectoDialog, "__init__", _prellenado)
    monkeypatch.setattr(CrearProyectoDialog, "exec", lambda self: 1)
    ventana.asistente.etapa_proyecto_perfil._on_crear_proyecto()


def _marcar_todos_los_documentos(etapa_cerrar_lote) -> None:
    """Selecciona (tilda) todas las filas de la tabla de EtapaCerrarLote --
    desde 2026-10-03 (NEXOM, seleccion de documentos al cerrar un lote)
    la tabla empieza SIN NADA marcado, asi que las pruebas que antes
    dependian de "se agregan todos los documentos automaticamente" ahora
    deben marcarlos explicitamente, igual que haria un operador real."""
    tabla = etapa_cerrar_lote.tabla_documentos
    for fila in range(tabla.rowCount()):
        tabla.item(fila, etapa_cerrar_lote._COL_SELECCION).setCheckState(Qt.CheckState.Checked)


def test_asistente_es_la_vista_predeterminada_al_abrir_la_ventana(qtbot, raiz_temporal) -> None:
    raiz, _origenes_dir = raiz_temporal
    ctx = bootstrap(raiz=str(raiz))
    ventana = MainWindow(ctx)
    qtbot.addWidget(ventana)
    # Desde la Fase 11.3, la pantalla principal es la vista de entrada
    # (indice 0); el asistente pasa a ser la vista predeterminada recien
    # AL ABRIR UN PROYECTO (seccion 4.11.1) -- ver
    # test_pantalla_principal_extremo_a_extremo.py para esa transicion.
    # Aqui solo se confirma que el asistente ya esta listo desde el
    # arranque (etapa 1, sin proyecto).
    assert ventana.radio_inicio.isChecked()
    assert ventana.asistente.etiqueta_progreso.text().startswith("Etapa 1 de 9")
    ventana.radio_asistente.setChecked(True)
    ventana._on_cambiar_modo(-1)
    assert ventana.modo_stack.currentWidget() is ventana.asistente


def test_recorrido_completo_de_un_documento_por_el_asistente(qtbot, raiz_temporal, monkeypatch) -> None:
    raiz, origenes_dir = raiz_temporal
    ctx = bootstrap(raiz=str(raiz))
    # Se registran las llamadas (en vez de un no-op ciego) para poder
    # verificar, mas abajo, que el cierre de lote exitoso usa
    # especificamente QMessageBox.information (gap de cobertura real
    # detectado en la auditoria 2026-07-31: antes se parchaban ambos como
    # no-op sin distinguir cual se habia llamado).
    llamadas_messagebox: list[tuple[str, tuple, dict]] = []
    monkeypatch.setattr(
        QMessageBox, "information", lambda *a, **k: llamadas_messagebox.append(("information", a, k))
    )
    monkeypatch.setattr(QMessageBox, "warning", lambda *a, **k: llamadas_messagebox.append(("warning", a, k)))
    # Confirmacion previa (2026-10-03, vista previa antes de cerrar el
    # lote): se acepta siempre en esta prueba, igual que los demas
    # QMessageBox parcheados arriba.
    monkeypatch.setattr(QMessageBox, "question", lambda *a, **k: QMessageBox.StandardButton.Yes)

    ventana = MainWindow(ctx)
    qtbot.addWidget(ventana)

    # -- Etapa 1: crear proyecto y elegir perfil (sin proyecto abierto todavia).
    assert ventana.proyecto_actual is None
    assert ventana.asistente.etapa_proyecto_perfil.etiqueta_estado.text() == "Todavia no hay ningun proyecto abierto."
    _crear_proyecto_desde_asistente(ventana, monkeypatch, "Asistente E2E")
    assert ventana.proyecto_actual is not None
    assert ventana.proyecto_actual.nombre == "Asistente E2E"
    assert ventana.asistente._etapa_actual == 0  # proyecto nuevo -- sin avance autoguardado todavia.

    # -- Etapa 2: importar y completar inventario (InventarioPanel embebido).
    ventana.asistente._mostrar_etapa(1)
    assert ventana.asistente.paginas.currentWidget() is ventana.asistente.etapa_inventario
    archivo = origenes_dir / "politica.docx"
    d = docx.Document()
    d.add_heading("Politica de prueba", level=1)
    d.add_paragraph("Contenido real para el asistente.")
    d.save(str(archivo))
    monkeypatch.setattr(QFileDialog, "getOpenFileNames", staticmethod(lambda *a, **k: ([str(archivo)], "")))
    ventana.asistente.etapa_inventario._on_agregar_documentos()
    assert ventana.asistente.etapa_inventario.tabla.rowCount() == 1
    documento = ventana.pctx.documents.listar_documentos()[0]

    # -- Etapa 3: procesar (clasificar + extraer, batch).
    ventana.asistente._mostrar_etapa(2)
    assert ventana.asistente.etapa_procesar.boton_procesar.isEnabled()
    ventana.asistente.etapa_procesar._on_procesar()
    documento = ventana.pctx.documents.obtener_documento(documento.id)
    assert documento.estado == EstadoDocumento.EXTRAIDO

    # -- Etapa 4: revisar -- deep-link al modo avanzado (Revision), donde
    # se completa de verdad (Modulo 8, sin cambios).
    ventana.asistente._mostrar_etapa(3)
    assert "1 documento(s)" in ventana.asistente.etapa_revisar.etiqueta_conteo.text()
    ventana.asistente.etapa_revisar.boton_ir.click()
    assert ventana.radio_modo_avanzado.isChecked()
    assert ventana.lista_navegacion.currentItem().text() == "Revision"
    ventana._cargar_documentos()
    ventana.tabla_documentos.selectRow(0)
    assert ventana.documento_actual.id == documento.id
    ventana.pctx.documents.revisar_documento(documento.id, "ana@nexom.mx", "ok", [])
    ventana.radio_asistente.setChecked(True)  # el operador vuelve al asistente.
    ventana._on_cambiar_modo(-1)

    # -- Etapa 5: resolver incidencias (IncidenciasPanel embebido, sin
    # incidencias pendientes en este recorrido -- solo se confirma que
    # la pantalla es real y esta vacia).
    ventana.asistente._mostrar_etapa(4)
    assert ventana.asistente.paginas.currentWidget() is ventana.asistente.etapa_incidencias
    assert ventana.asistente.etapa_incidencias.tabla.rowCount() == 0

    # -- Etapa 6: normalizar (batch).
    ventana.asistente._mostrar_etapa(5)
    ventana.asistente.etapa_normalizar.establecer_proyecto(ventana.pctx, ventana.proyecto_actual)
    assert ventana.asistente.etapa_normalizar.boton_normalizar.isEnabled()
    ventana.asistente.etapa_normalizar._on_normalizar()
    documento = ventana.pctx.documents.obtener_documento(documento.id)
    assert documento.estado == EstadoDocumento.NORMALIZADO

    # -- Etapa 7: aprobar (batch, con aprobador comun).
    ventana.asistente._mostrar_etapa(6)
    ventana.asistente.etapa_aprobar.establecer_proyecto(ventana.pctx, ventana.proyecto_actual)
    ventana.asistente.etapa_aprobar.campo_aprobador.setText("ana@nexom.mx")
    # Bug real corregido 2026-07-31 (tercera auditoria general): esta
    # etapa descartaba el comentario de aprobacion (pasaba "" fijo) --
    # se confirma aqui que el campo nuevo si se propaga hasta la
    # Aprobacion persistida, igual que en Modo Avanzado.
    ventana.asistente.etapa_aprobar.campo_comentario.setText("Cumple los criterios de calidad.")
    ventana.asistente.etapa_aprobar._on_aprobar()
    documento = ventana.pctx.documents.obtener_documento(documento.id)
    assert documento.estado == EstadoDocumento.APROBADO_PARA_MARKDOWN
    aprobaciones = ventana.pctx.documents.listar_aprobaciones(documento.id)
    assert aprobaciones[-1].comentario == "Cumple los criterios de calidad."

    # -- Etapa 8: generar Markdown (batch: generar + validar).
    ventana.asistente._mostrar_etapa(7)
    ventana.asistente.etapa_markdown.establecer_proyecto(ventana.pctx, ventana.proyecto_actual)
    ventana.asistente.etapa_markdown._on_generar()
    documento = ventana.pctx.documents.obtener_documento(documento.id)
    assert documento.estado == EstadoDocumento.VALIDADO

    # Fuente requerida antes de poder liberar (Modulo 11/13, ver
    # test_gui_menu_completo_extremo_a_extremo.py para el mismo hallazgo).
    ventana.pctx.validity.registrar_fuente(documento.id, "sistema de prueba", "Direccion Juridica", "confirmado")

    # "Ejecutar pruebas" ya no es una etapa del asistente (2026-07-31,
    # solicitud explicita del usuario): PruebasPanel es una herramienta
    # de QA que valida el motor de extraccion/OCR en general contra un
    # corpus ad-hoc, no un paso del flujo por proyecto -- sigue
    # disponible de forma independiente en Modo Avanzado
    # ("Pruebas de aceptacion").

    # -- Etapa 9: cerrar lote (crea el lote, agrega todo, cierra --
    # libera lo que ya tiene fuente registrada).
    ventana.asistente._mostrar_etapa(8)
    ventana.asistente.etapa_cerrar_lote.establecer_proyecto(ventana.pctx, ventana.proyecto_actual)
    assert ventana.asistente.etapa_cerrar_lote.campo_nombre.text() == f"Lote {ventana.proyecto_actual.nombre}"
    _marcar_todos_los_documentos(ventana.asistente.etapa_cerrar_lote)
    ventana.asistente.etapa_cerrar_lote._on_cerrar_lote()
    documento = ventana.pctx.documents.obtener_documento(documento.id)
    assert documento.estado == EstadoDocumento.LIBERADO
    assert "1 liberado(s), 0 bloqueado(s)" in ventana.asistente.etapa_cerrar_lote.etiqueta_resultado.text()

    # La entrega a RAG Studio se preparo sin problema -- se avisa con
    # QMessageBox.information (no .warning), gap de cobertura cerrado en
    # esta misma pasada (ver tambien
    # test_cierre_de_lote_avisa_con_advertencia_si_falla_la_entrega_a_rag_studio
    # para el camino contrario).
    assert llamadas_messagebox, "se esperaba al menos una llamada a QMessageBox durante el cierre de lote"
    tipo_ultima_llamada, args_ultima_llamada, _ = llamadas_messagebox[-1]
    assert tipo_ultima_llamada == "information"
    assert args_ultima_llamada[1] == "Lote cerrado"

    ventana.asistente.etapa_cerrar_lote.boton_ir_reportes.click()
    assert ventana.radio_modo_avanzado.isChecked()
    assert ventana.lista_navegacion.currentItem().text() == "Reportes"


def test_avance_del_asistente_se_recupera_al_reabrir_el_proyecto(qtbot, raiz_temporal, monkeypatch) -> None:
    raiz, _origenes_dir = raiz_temporal
    ctx = bootstrap(raiz=str(raiz))
    proyecto = ctx.projects.crear_proyecto(nombre="Recuperacion", perfil_id="perfil_generico")

    ventana1 = MainWindow(ctx)
    qtbot.addWidget(ventana1)
    ventana1.combo_proyectos.setCurrentIndex(ventana1.combo_proyectos.findData(proyecto.id))
    ventana1._on_abrir_proyecto()
    ventana1.asistente._mostrar_etapa(4)  # "Resolver incidencias"

    archivo_estado = Path(proyecto.ruta_carpeta) / ".asistente_estado.json"
    assert archivo_estado.exists()
    assert '"etapa": 4' in archivo_estado.read_text(encoding="utf-8")

    # Una ventana nueva simula reabrir la aplicacion tras un cierre
    # inesperado (seccion 4.11.1: "sesiones recuperables").
    ventana2 = MainWindow(ctx)
    qtbot.addWidget(ventana2)
    ventana2.combo_proyectos.setCurrentIndex(ventana2.combo_proyectos.findData(proyecto.id))
    ventana2._on_abrir_proyecto()
    assert ventana2.asistente._etapa_actual == 4
    assert ventana2.asistente.paginas.currentWidget() is ventana2.asistente.etapa_incidencias


def test_no_se_puede_saltar_a_una_etapa_no_alcanzada(qtbot, raiz_temporal) -> None:
    raiz, _origenes_dir = raiz_temporal
    ctx = bootstrap(raiz=str(raiz))
    ventana = MainWindow(ctx)
    qtbot.addWidget(ventana)

    assert ventana.asistente._etapa_actual == 0
    ventana.asistente.lista_etapas.setCurrentRow(5)  # todavia no alcanzada.
    assert ventana.asistente._etapa_actual == 0
    assert ventana.asistente.lista_etapas.currentRow() == 0

    ventana.asistente._mostrar_etapa(3)  # avanza "de verdad" (simula haber llegado hasta ahi).
    ventana.asistente.lista_etapas.setCurrentRow(1)  # retroceder a una ya alcanzada si es valido.
    assert ventana.asistente._etapa_actual == 1
    ventana.asistente.lista_etapas.setCurrentRow(3)
    assert ventana.asistente._etapa_actual == 3
    ventana.asistente.lista_etapas.setCurrentRow(8)  # todavia no alcanzada.
    assert ventana.asistente._etapa_actual == 3


def test_cierre_de_lote_avisa_con_advertencia_si_falla_la_entrega_a_rag_studio(
    qtbot, raiz_temporal, monkeypatch
) -> None:
    """Gap de cobertura real detectado en la auditoria 2026-07-31: el
    QMessageBox agregado a EtapaCerrarLote._on_cerrar_lote() distingue
    entre 'Lote cerrado' (QMessageBox.information, camino feliz -- ya
    cubierto en test_recorrido_completo_de_un_documento_por_el_asistente)
    y 'Lote cerrado con avisos' (QMessageBox.warning, cuando el lote se
    cierra bien pero preparar_entrega_rag_studio() falla, por ejemplo por
    permisos de escritura). Se fuerza esa segunda rama monkeypencheando
    el metodo del servicio para que lance, sin tocar nada del sistema de
    archivos real."""
    raiz, origenes_dir = raiz_temporal
    ctx = bootstrap(raiz=str(raiz))
    llamadas_messagebox: list[tuple[str, tuple, dict]] = []
    monkeypatch.setattr(
        QMessageBox, "information", lambda *a, **k: llamadas_messagebox.append(("information", a, k))
    )
    monkeypatch.setattr(QMessageBox, "warning", lambda *a, **k: llamadas_messagebox.append(("warning", a, k)))
    monkeypatch.setattr(QMessageBox, "question", lambda *a, **k: QMessageBox.StandardButton.Yes)

    ventana = MainWindow(ctx)
    qtbot.addWidget(ventana)
    _crear_proyecto_desde_asistente(ventana, monkeypatch, "Asistente entrega con falla")
    assert ventana.pctx is not None
    assert ventana.proyecto_actual is not None

    # Un documento inventariado basta para que el lote no quede vacio
    # (test_cerrar_lote_vacio_falla confirma que un lote sin documentos no
    # se puede cerrar) -- no hace falta llevarlo hasta LIBERADO, solo que
    # el cierre del lote en si mismo tenga exito para llegar a la
    # preparacion de la entrega.
    archivo = origenes_dir / "memo.txt"
    archivo.write_text("Contenido real para la prueba de la advertencia.", encoding="utf-8")
    ventana.pctx.documents.registrar_documento(proyecto_id=ventana.proyecto_actual.id, ruta_origen=str(archivo))

    monkeypatch.setattr(
        ventana.pctx.documents,
        "preparar_entrega_rag_studio",
        lambda lote_id=None: (_ for _ in ()).throw(RuntimeError("fallo simulado de permisos")),
    )

    ventana.asistente._mostrar_etapa(8)
    ventana.asistente.etapa_cerrar_lote.establecer_proyecto(ventana.pctx, ventana.proyecto_actual)
    _marcar_todos_los_documentos(ventana.asistente.etapa_cerrar_lote)
    ventana.asistente.etapa_cerrar_lote._on_cerrar_lote()

    assert llamadas_messagebox, "se esperaba al menos una llamada a QMessageBox durante el cierre de lote"
    tipo_ultima_llamada, args_ultima_llamada, _ = llamadas_messagebox[-1]
    assert tipo_ultima_llamada == "warning"
    assert args_ultima_llamada[1] == "Lote cerrado con avisos"
    assert "No se pudo preparar la entrega para RAG Studio" in ventana.asistente.etapa_cerrar_lote.etiqueta_resultado.text()
