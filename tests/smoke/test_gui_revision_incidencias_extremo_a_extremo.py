"""Prueba de humo del prototipo de GUI (Modulo 8 + panel de Modulo 10).
Usa una raiz temporal, nunca la ruta real de datos -- mismo patron que
el resto de tests/smoke/. Requiere PySide6 y pytest-qt (extras `ui` y
`ui-dev`, ver pyproject.toml) y, en este entorno de desarrollo sin
sistema grafico real, las bibliotecas EGL/OpenGL extraidas manualmente
(ver README, "Interfaz grafica (prototipo)": desviacion de entorno) mas
QT_QPA_PLATFORM=offscreen.

Los dialogos modales (RevisarDocumentoDialog/RegistrarIncidenciaDialog)
se parchan para no bloquear la prueba en modo headless -- se les
inyectan valores y se simula "Aceptar" sin abrir ninguna ventana real,
igual que haria un usuario llenando el formulario.
"""
from __future__ import annotations

import shutil
from pathlib import Path

import docx
import pytest

pytest.importorskip("PySide6")

from PySide6.QtWidgets import QMessageBox  # noqa: E402

from app.bootstrap import bootstrap  # noqa: E402
from domain.documents.entities import Criticidad, Documento, EstadoDocumento  # noqa: E402
from domain.documents.ocr_table_reconstruction import RangoBajaConfianza  # noqa: E402
from domain.incidents.entities import EtapaIncidencia, SeveridadIncidencia  # noqa: E402
from ui.dialogs import RegistrarIncidenciaDialog, RegistrarVigenciaDialog, RevisarDocumentoDialog  # noqa: E402
from ui.main_window import MainWindow  # noqa: E402

RAIZ_PRUEBA = Path("/tmp/rag_content_studio_smoke_test_gui")


@pytest.fixture()
def raiz_temporal():
    if RAIZ_PRUEBA.exists():
        shutil.rmtree(RAIZ_PRUEBA)
    RAIZ_PRUEBA.mkdir(parents=True)
    perfiles_dir = RAIZ_PRUEBA / "perfiles" / "perfil_generico"
    perfiles_dir.mkdir(parents=True)
    shutil.copy(
        Path(__file__).resolve().parents[2] / "perfiles" / "perfil_generico" / "0.1.yaml",
        perfiles_dir / "0.1.yaml",
    )
    origenes_dir = RAIZ_PRUEBA / "_origenes_de_prueba"
    origenes_dir.mkdir()
    yield RAIZ_PRUEBA, origenes_dir
    shutil.rmtree(RAIZ_PRUEBA, ignore_errors=True)


def _ir_a_revision(ventana: MainWindow) -> None:
    """Navega a la pestana "Revision" igual que lo haria un operador
    real (clic en la lista de navegacion), sin ninguna llamada manual a
    `ventana._cargar_documentos()`. Bug real corregido 2026-07-30: antes
    de esa fecha, esta navegacion no refrescaba `tabla_documentos` --
    todas las pruebas de este archivo necesitaban forzar el refresco a
    mano porque el documento se registraba/procesaba DESPUES de abrir
    el proyecto (mismo escenario que un operador procesando en otra
    pestana mientras "Revision" ya estaba cargada)."""
    indice = next(i for i in range(ventana.lista_navegacion.count()) if ventana.lista_navegacion.item(i).text() == "Revision")
    ventana.lista_navegacion.setCurrentRow(indice)


def test_revisar_y_registrar_incidencia_desde_la_ventana_principal(qtbot, raiz_temporal, monkeypatch) -> None:
    raiz, origenes_dir = raiz_temporal
    ctx = bootstrap(raiz=str(raiz))
    proyecto = ctx.projects.crear_proyecto(nombre="GUI", perfil_id="perfil_generico")

    archivo = origenes_dir / "contrato.docx"
    d = docx.Document()
    d.add_heading("Contrato de prueba", level=1)
    d.add_paragraph("Contenido de prueba para la GUI.")
    d.save(str(archivo))

    ventana = MainWindow(ctx)
    qtbot.addWidget(ventana)

    assert ventana.combo_proyectos.count() == 1
    ventana._on_abrir_proyecto()
    assert ventana.pctx is not None

    documento = ventana.pctx.documents.registrar_documento(proyecto_id=proyecto.id, ruta_origen=str(archivo))
    ventana.pctx.documents.clasificar_documento(documento.id)
    extraido = ventana.pctx.documents.extraer_documento(documento.id)

    _ir_a_revision(ventana)
    assert ventana.tabla_documentos.rowCount() == 1
    ventana.tabla_documentos.selectRow(0)

    assert ventana.documento_actual is not None
    assert ventana.documento_actual.id == extraido.id
    assert "Contenido de prueba para la GUI." in ventana.texto_extraido.toPlainText()
    assert ventana.boton_revisar.isEnabled()

    # -- Aprobar revision: se parcha el dialogo modal para que ya
    # tenga los valores capturados y "Aceptar" simulado, sin bloquear
    # la prueba con una ventana real.
    def _dialogo_revision_prellenado(self, parent=None):
        RevisarDocumentoDialog.__base_init__(self, parent)
        self.campo_revisor.setText("ana@nexom.mx")
        self.campo_comentario.setPlainText("Cotejado contra el original, todo correcto.")

    RevisarDocumentoDialog.__base_init__ = RevisarDocumentoDialog.__init__
    monkeypatch.setattr(RevisarDocumentoDialog, "__init__", _dialogo_revision_prellenado)
    monkeypatch.setattr(RevisarDocumentoDialog, "exec", lambda self: 1)

    ventana._on_aprobar_revision()
    assert ventana.documento_actual.estado == EstadoDocumento.REVISADO
    assert ventana.boton_revisar.isEnabled() is False

    # Modulo 8: la revision queda en su propio historial (RevisionRepository),
    # no solo en el evento de auditoria -- en este momento el documento
    # todavia no tiene ninguna incidencia registrada.
    revisiones = ventana.pctx.documents.listar_revisiones(ventana.documento_actual.id)
    assert len(revisiones) == 1
    assert revisiones[0].revisor == "ana@nexom.mx"
    assert revisiones[0].incidencias_detectadas == []

    # -- Registrar incidencia: mismo patron de parche que arriba.
    def _dialogo_incidencia_prellenado(self, parent=None):
        RegistrarIncidenciaDialog.__base_init__(self, parent)
        indice_etapa = self.campo_etapa.findData(EtapaIncidencia.REVISION)
        self.campo_etapa.setCurrentIndex(indice_etapa)
        indice_severidad = self.campo_severidad.findData(SeveridadIncidencia.CRITICA)
        self.campo_severidad.setCurrentIndex(indice_severidad)
        self.campo_tipo.setText("cifra_dudosa")
        self.campo_evidencia.setPlainText("La cifra de la pagina 1 no coincide con el original.")
        self.campo_impacto.setPlainText("Podria invalidar el monto reportado.")

    RegistrarIncidenciaDialog.__base_init__ = RegistrarIncidenciaDialog.__init__
    monkeypatch.setattr(RegistrarIncidenciaDialog, "__init__", _dialogo_incidencia_prellenado)
    monkeypatch.setattr(RegistrarIncidenciaDialog, "exec", lambda self: 1)

    ventana._on_registrar_incidencia()

    incidencias = ventana.pctx.incidents.listar_incidencias_de_documento(ventana.documento_actual.id)
    assert len(incidencias) == 1
    assert incidencias[0].severidad == SeveridadIncidencia.CRITICA
    assert ventana.lista_incidencias.count() == 1
    assert "criticas abiertas" in ventana.etiqueta_estado_documento.text().lower()


def test_paneles_de_documentos_y_estado_se_pueden_ocultar_y_mostrar(qtbot, raiz_temporal) -> None:
    """Bug real reportado 2026-07-31 (captura de pantalla del consultor):
    el operador queria poder ocultar los paneles izquierdo ("Documentos
    del proyecto") y derecho ("Estado") de la pantalla Revision para
    aprovechar mejor el visor central Original/Normalizado -- ver
    docstring de MainWindow._construir_pagina_revision().

    Bug real del propio test, encontrado 2026-08-01 (la falla solo salio
    a la luz al poder correr por fin `tests/smoke/` de verdad en el
    entorno de trabajo, en vez de solo verificarlo con `py_compile`): las
    aserciones de `isVisible()` de abajo dependen de DOS cosas que este
    test no hacia. Primero, `ventana.show()` -- sin exponer la ventana de
    verdad, `isVisible()` de cualquier widget hijo es `False` por
    semantica normal de Qt, sin relacion con el bug que se queria probar.
    Segundo, entrar de verdad a "Modo avanzado": `_ir_a_revision()` solo
    cambia `lista_navegacion`, que vive DENTRO de la pagina "Modo
    avanzado" de `modo_stack` (indice 2) -- pero `_on_abrir_proyecto()`
    (Fase 11.3) navega automaticamente a "Asistente" (indice 1) al abrir
    un proyecto desde "Inicio", asi que sin cambiar de modo explicitamente
    la pagina "Revision" quedaba seleccionada dentro de un stack que
    nunca era el visible. `radio_modo_avanzado.setChecked(True)` no
    alcanza para eso: `MainWindow._on_cambiar_modo` esta conectado a
    `grupo_modos.idClicked`, una señal que solo se dispara con un clic
    real, no al fijar el estado del boton a mano -- de ahi `.click()`
    abajo. El feature en si (ocultar/mostrar los paneles) nunca estuvo
    roto: un script de reproduccion manual con `ventana.show()` +
    `radio_modo_avanzado.click()` + `qtbot`/`processEvents()` confirmo
    que ambos paneles se ocultan y muestran correctamente; solo la
    asercion de este test no reflejaba una navegacion real."""
    raiz, _origenes_dir = raiz_temporal
    ctx = bootstrap(raiz=str(raiz))
    ctx.projects.crear_proyecto(nombre="GUI", perfil_id="perfil_generico")

    ventana = MainWindow(ctx)
    qtbot.addWidget(ventana)
    with qtbot.waitExposed(ventana):
        ventana.show()
    ventana._on_abrir_proyecto()
    ventana.radio_modo_avanzado.click()  # entrar de verdad a "Modo avanzado", ver docstring arriba.
    _ir_a_revision(ventana)

    # Por defecto, ambos paneles estan visibles.
    assert ventana._panel_documentos_revision.isVisible()
    assert ventana._area_lateral_revision.isVisible()
    assert ventana.boton_alternar_panel_documentos.text() == "Ocultar documentos"
    assert ventana.boton_alternar_panel_estado.text() == "Ocultar estado"

    ventana.boton_alternar_panel_documentos.setChecked(False)
    assert not ventana._panel_documentos_revision.isVisible()
    assert ventana.boton_alternar_panel_documentos.text() == "Mostrar documentos"
    assert ventana._area_lateral_revision.isVisible()  # el otro panel no se ve afectado

    ventana.boton_alternar_panel_estado.setChecked(False)
    assert not ventana._area_lateral_revision.isVisible()
    assert ventana.boton_alternar_panel_estado.text() == "Mostrar estado"

    # Se pueden volver a mostrar.
    ventana.boton_alternar_panel_documentos.setChecked(True)
    ventana.boton_alternar_panel_estado.setChecked(True)
    assert ventana._panel_documentos_revision.isVisible()
    assert ventana._area_lateral_revision.isVisible()
    assert ventana.boton_alternar_panel_documentos.text() == "Ocultar documentos"
    assert ventana.boton_alternar_panel_estado.text() == "Ocultar estado"


def test_no_se_puede_aprobar_revision_con_incidencia_critica_abierta(qtbot, raiz_temporal, monkeypatch) -> None:
    # Modulo 13 (gobernanza): documenta el bloqueo real, no solo el
    # aviso -- una incidencia critica registrada ANTES de intentar
    # aprobar debe deshabilitar el boton (UX proactiva) y, si se
    # invoca el manejador de todas formas, GovernanceService debe
    # impedir la transicion de verdad (ver
    # application.governance.service.GovernanceService).
    raiz, origenes_dir = raiz_temporal
    ctx = bootstrap(raiz=str(raiz))
    proyecto = ctx.projects.crear_proyecto(nombre="GUI-bloqueo", perfil_id="perfil_generico")

    archivo = origenes_dir / "contrato.docx"
    d = docx.Document()
    d.add_heading("Contrato de prueba", level=1)
    d.add_paragraph("Contenido de prueba para la GUI.")
    d.save(str(archivo))

    ventana = MainWindow(ctx)
    qtbot.addWidget(ventana)
    ventana._on_abrir_proyecto()

    documento = ventana.pctx.documents.registrar_documento(proyecto_id=proyecto.id, ruta_origen=str(archivo))
    ventana.pctx.documents.clasificar_documento(documento.id)
    extraido = ventana.pctx.documents.extraer_documento(documento.id)

    ventana.pctx.incidents.registrar_incidencia(
        proyecto_id=proyecto.id,
        documento_id=extraido.id,
        etapa=EtapaIncidencia.REVISION,
        tipo="cifra_dudosa",
        severidad=SeveridadIncidencia.CRITICA,
        evidencia="La cifra de la pagina 1 no coincide con el original.",
        impacto="Podria invalidar el monto reportado.",
    )

    _ir_a_revision(ventana)
    ventana.tabla_documentos.selectRow(0)
    assert ventana.documento_actual.id == extraido.id
    # UX proactiva: el boton ya esta deshabilitado antes de intentar nada.
    assert ventana.boton_revisar.isEnabled() is False
    assert "criticas abiertas" in ventana.etiqueta_estado_documento.text().lower()

    def _dialogo_revision_prellenado(self, parent=None):
        RevisarDocumentoDialog.__base_init__(self, parent)
        self.campo_revisor.setText("ana@nexom.mx")

    RevisarDocumentoDialog.__base_init__ = RevisarDocumentoDialog.__init__
    monkeypatch.setattr(RevisarDocumentoDialog, "__init__", _dialogo_revision_prellenado)
    monkeypatch.setattr(RevisarDocumentoDialog, "exec", lambda self: 1)
    monkeypatch.setattr(QMessageBox, "critical", lambda *args, **kwargs: None)

    # Se invoca el manejador directamente (bypass del boton deshabilitado,
    # como podria pasar si algo mas lo dispara) para confirmar que el
    # bloqueo real vive en application, no solo en el estado del boton.
    ventana._on_aprobar_revision()

    assert ventana.pctx.documents.obtener_documento(extraido.id).estado == EstadoDocumento.EXTRAIDO
    assert ventana.pctx.documents.listar_revisiones(extraido.id) == []


def test_no_se_puede_aprobar_revision_de_documento_critico_con_vigencia_desconocida(
    qtbot, raiz_temporal, monkeypatch
) -> None:
    # Modulo 13 (gobernanza), segunda regla (2026-07-28): un documento
    # CRITICO sin vigencia registrada debe deshabilitar el boton
    # (UX proactiva) y, si se invoca el manejador de todas formas,
    # GovernanceService debe impedir la transicion de verdad -- mismo
    # criterio que la primera regla (incidencias criticas abiertas),
    # arriba.
    raiz, origenes_dir = raiz_temporal
    ctx = bootstrap(raiz=str(raiz))
    proyecto = ctx.projects.crear_proyecto(nombre="GUI-vigencia-bloqueo", perfil_id="perfil_generico")

    archivo = origenes_dir / "contrato.docx"
    d = docx.Document()
    d.add_heading("Contrato critico de prueba", level=1)
    d.add_paragraph("Contenido de prueba para la GUI.")
    d.save(str(archivo))

    ventana = MainWindow(ctx)
    qtbot.addWidget(ventana)
    ventana._on_abrir_proyecto()

    documento = ventana.pctx.documents.registrar_documento(
        proyecto_id=proyecto.id, ruta_origen=str(archivo), criticidad=Criticidad.CRITICO
    )
    ventana.pctx.documents.clasificar_documento(documento.id)
    extraido = ventana.pctx.documents.extraer_documento(documento.id)

    _ir_a_revision(ventana)
    ventana.tabla_documentos.selectRow(0)
    assert ventana.documento_actual.id == extraido.id
    # UX proactiva: el boton ya esta deshabilitado antes de intentar nada,
    # sin que se haya registrado ninguna incidencia critica.
    assert ventana.boton_revisar.isEnabled() is False
    assert "vigencia es desconocida" in ventana.etiqueta_estado_documento.text().lower()

    def _dialogo_revision_prellenado(self, parent=None):
        RevisarDocumentoDialog.__base_init__(self, parent)
        self.campo_revisor.setText("ana@nexom.mx")

    RevisarDocumentoDialog.__base_init__ = RevisarDocumentoDialog.__init__
    monkeypatch.setattr(RevisarDocumentoDialog, "__init__", _dialogo_revision_prellenado)
    monkeypatch.setattr(RevisarDocumentoDialog, "exec", lambda self: 1)
    monkeypatch.setattr(QMessageBox, "critical", lambda *args, **kwargs: None)

    # Se invoca el manejador directamente (bypass del boton deshabilitado)
    # para confirmar que el bloqueo real vive en application, no solo en
    # el estado del boton.
    ventana._on_aprobar_revision()

    assert ventana.pctx.documents.obtener_documento(extraido.id).estado == EstadoDocumento.EXTRAIDO
    assert ventana.pctx.documents.listar_revisiones(extraido.id) == []

    # Registrar la vigencia (completa, sin vencimiento) debe habilitar
    # el boton de nuevo.
    def _dialogo_vigencia_prellenado(self, parent=None, fuente=None, vigencia=None, documentos_disponibles=None):
        RegistrarVigenciaDialog.__base_init__(self, parent, fuente, vigencia, documentos_disponibles)
        self.campo_sistema_origen.setText("portal de normatividad interna")
        self.campo_autoridad_fuente.setText("Direccion Juridica")
        self.campo_evidencia_consulta.setText("Confirmado el 2026-07-28 con el area emisora.")
        self.campo_vigente_desde.setText("2026-01-01")
        self.campo_sin_vencimiento.setChecked(True)
        self.campo_autoridad_vigencia.setText("Direccion Juridica")

    RegistrarVigenciaDialog.__base_init__ = RegistrarVigenciaDialog.__init__
    monkeypatch.setattr(RegistrarVigenciaDialog, "__init__", _dialogo_vigencia_prellenado)
    monkeypatch.setattr(RegistrarVigenciaDialog, "exec", lambda self: 1)

    ventana._on_registrar_vigencia_fuente()

    assert ventana.boton_revisar.isEnabled() is True
    assert "vigencia es desconocida" not in ventana.etiqueta_estado_documento.text().lower()

    ventana._on_aprobar_revision()
    assert ventana.documento_actual.estado == EstadoDocumento.REVISADO


def test_correccion_verificaciones_y_navegacion_sincronizada(qtbot, raiz_temporal, monkeypatch) -> None:
    """Modulo 8, cierre de los tres pendientes del primer prototipo:
    correccion editable persistida, deteccion de fechas/cifras, y
    navegacion sincronizada real usando mapa_paginas (solo lo llena
    PdfExtractor en un PDF de mas de una pagina, ver su docstring)."""
    from reportlab.pdfgen import canvas

    raiz, origenes_dir = raiz_temporal
    ctx = bootstrap(raiz=str(raiz))
    proyecto = ctx.projects.crear_proyecto(nombre="GUI-correccion", perfil_id="perfil_generico")

    archivo = origenes_dir / "contrato.pdf"
    c = canvas.Canvas(str(archivo))
    c.drawString(100, 700, "Primera pagina, vigente desde 12/05/2026.")
    c.showPage()
    c.drawString(100, 700, "Segunda pagina, monto total $1,234.56.")
    c.save()

    ventana = MainWindow(ctx)
    qtbot.addWidget(ventana)
    ventana._on_abrir_proyecto()

    documento = ventana.pctx.documents.registrar_documento(proyecto_id=proyecto.id, ruta_origen=str(archivo))
    ventana.pctx.documents.clasificar_documento(documento.id)
    extraido = ventana.pctx.documents.extraer_documento(documento.id)
    assert len(extraido.mapa_paginas) == 2  # precondicion: si esto falla, no es un bug de la GUI

    _ir_a_revision(ventana)
    ventana.tabla_documentos.selectRow(0)
    assert ventana.documento_actual.id == extraido.id

    # -- Verificacion de fechas/cifras: ambos hallazgos deben listarse.
    textos_detectados = [ventana.lista_verificaciones.item(i).text() for i in range(ventana.lista_verificaciones.count())]
    assert any("12/05/2026" in t for t in textos_detectados)
    assert any("$1,234.56" in t for t in textos_detectados)

    # -- Navegacion sincronizada: pagina 1 -> cursor al inicio del
    # markdown (offset 0); pagina 2 -> cursor donde empieza esa pagina.
    assert ventana.texto_extraido.textCursor().position() == 0
    ventana._pagina_pdf_siguiente()
    offset_pagina_2 = extraido.mapa_paginas[1]
    assert ventana.texto_extraido.textCursor().position() == offset_pagina_2
    assert ventana.texto_extraido.toPlainText()[offset_pagina_2:].startswith("Segunda pagina")

    # -- Correccion editable: documento EXTRAIDO permite editar y guardar.
    assert ventana.texto_extraido.isReadOnly() is False
    assert ventana.boton_guardar_correccion.isEnabled()
    ventana.texto_extraido.setPlainText("Contenido corregido a mano por el revisor.")
    ventana._on_guardar_correccion()

    corregido = ventana.pctx.documents.obtener_documento(extraido.id)
    assert corregido.ruta_revisado is not None
    assert corregido.estado == EstadoDocumento.EXTRAIDO  # corregir() no avanza de estado
    assert Path(corregido.ruta_revisado).read_text(encoding="utf-8") == "Contenido corregido a mano por el revisor."

    # -- Tras aprobar la revision, el panel deja de ser editable.
    def _dialogo_revision_prellenado(self, parent=None):
        RevisarDocumentoDialog.__base_init__(self, parent)
        self.campo_revisor.setText("ana@nexom.mx")

    RevisarDocumentoDialog.__base_init__ = RevisarDocumentoDialog.__init__
    monkeypatch.setattr(RevisarDocumentoDialog, "__init__", _dialogo_revision_prellenado)
    monkeypatch.setattr(RevisarDocumentoDialog, "exec", lambda self: 1)

    ventana._on_aprobar_revision()

    assert ventana.documento_actual.estado == EstadoDocumento.REVISADO
    assert ventana.texto_extraido.isReadOnly() is True
    assert ventana.boton_guardar_correccion.isEnabled() is False


def test_resaltado_de_baja_confianza_ocr_se_muestra_en_revision(qtbot, raiz_temporal) -> None:
    """Nivel 1 (confianza de OCR), 2026-08-01: las palabras que Tesseract
    marco con baja confianza (domain.documents.ocr_table_reconstruction.
    RangoBajaConfianza) deben resaltarse con fondo ambar en el panel
    "Extraido / revisado" y el resumen debe reflejar cuantas se
    resaltaron -- ver MainWindow._resaltar_baja_confianza(). No se usa
    OCR real aqui (los textos nitidos de las pruebas de
    test_tesseract_ocr_engine.py no producen baja confianza a proposito)
    -- se construye el Documento con rangos_baja_confianza ya fijados,
    igual que si vinieran de una extraccion OCR real, para poder probar
    el resaltado de la GUI de forma determinista."""
    raiz, origenes_dir = raiz_temporal
    ctx = bootstrap(raiz=str(raiz))
    proyecto = ctx.projects.crear_proyecto(nombre="GUI-confianza", perfil_id="perfil_generico")

    archivo = origenes_dir / "contrato.docx"
    d = docx.Document()
    d.add_heading("Encabezado", level=1)
    d.add_paragraph("Contenido de prueba para la GUI.")
    d.save(str(archivo))

    ventana = MainWindow(ctx)
    qtbot.addWidget(ventana)
    ventana._on_abrir_proyecto()

    documento = ventana.pctx.documents.registrar_documento(proyecto_id=proyecto.id, ruta_origen=str(archivo))
    ventana.pctx.documents.clasificar_documento(documento.id)
    extraido = ventana.pctx.documents.extraer_documento(documento.id)

    _ir_a_revision(ventana)
    ventana.tabla_documentos.selectRow(0)
    assert ventana.documento_actual.id == extraido.id
    # Extraccion directa (no OCR): sin rangos_baja_confianza, nada que
    # resaltar.
    assert ventana.texto_extraido.extraSelections() == []
    assert ventana.etiqueta_baja_confianza.text() == ""

    texto_actual = ventana.texto_extraido.toPlainText()
    inicio = texto_actual.index("prueba")
    fin = inicio + len("prueba")
    data = extraido.model_dump()
    data["rangos_baja_confianza"] = [RangoBajaConfianza(inicio=inicio, fin=fin, confianza=42.0, texto="prueba")]
    con_baja_confianza = Documento(**data)

    ventana.documento_actual = con_baja_confianza
    ventana._mostrar_extraido(con_baja_confianza)

    selecciones = ventana.texto_extraido.extraSelections()
    assert len(selecciones) == 1
    assert selecciones[0].cursor.selectedText() == "prueba"
    assert "1 palabra" in ventana.etiqueta_baja_confianza.text()

    # Rango fuera de los limites del texto vigente -- mismo criterio
    # aceptado que mapa_paginas (ver docstring de
    # `Documento.rangos_baja_confianza` y de
    # `MainWindow._resaltar_baja_confianza()`): se descarta en silencio,
    # nunca lanza.
    data["rangos_baja_confianza"] = [
        RangoBajaConfianza(inicio=10_000, fin=10_010, confianza=10.0, texto="fantasma")
    ]
    fuera_de_rango = Documento(**data)
    ventana.documento_actual = fuera_de_rango
    ventana._mostrar_extraido(fuera_de_rango)
    assert ventana.texto_extraido.extraSelections() == []
    assert ventana.etiqueta_baja_confianza.text() == ""


def test_verificaciones_de_nivel_2_se_resaltan_en_la_lista(qtbot, raiz_temporal) -> None:
    """Nivel 2 (validaciones numericas), 2026-08-01: una fecha imposible
    segun el calendario (domain.documents.verificacion_fechas_cifras) y
    una tabla cuyo total no cuadra con la suma de sus filas
    (domain.documents.verificacion_totales_tabla) deben aparecer
    combinados en `lista_verificaciones`, ordenados por posicion, con un
    marcador visual (mismo ambar que Nivel 1) y un tooltip con el motivo
    -- ver MainWindow._mostrar_verificaciones(). Se llama directamente
    con un Markdown fabricado (en vez de extraer un documento real) para
    aislar el comportamiento de la GUI del de la deteccion en si, que ya
    tiene su propia cobertura exhaustiva en
    tests/unit/domain/test_verificacion_fechas_cifras.py y
    tests/unit/domain/test_verificacion_totales_tabla.py."""
    from PySide6.QtGui import QColor

    from ui.theme import AMBAR_ATENCION_FONDO

    raiz, _origenes_dir = raiz_temporal
    ctx = bootstrap(raiz=str(raiz))
    ctx.projects.crear_proyecto(nombre="GUI-nivel2", perfil_id="perfil_generico")

    ventana = MainWindow(ctx)
    qtbot.addWidget(ventana)

    markdown = "\n".join(
        [
            "Contrato vigente desde el 31/04/2026 (fecha invalida a proposito).",
            "",
            "| Producto | Cantidad |",
            "| --- | --- |",
            "| A | 10 |",
            "| B | 20 |",
            "| Total | 999 |",
        ]
    )
    ventana._mostrar_verificaciones(markdown)

    assert ventana.lista_verificaciones.count() == 2
    item_fecha = ventana.lista_verificaciones.item(0)
    item_total = ventana.lista_verificaciones.item(1)

    assert "31/04/2026" in item_fecha.text()
    assert item_fecha.text().startswith("⚠")
    assert item_fecha.toolTip() != ""
    assert item_fecha.background().color() == QColor(AMBAR_ATENCION_FONDO)

    assert "999" in item_total.text()
    assert item_total.text().startswith("⚠")
    assert "30" in item_total.toolTip()
    assert item_total.background().color() == QColor(AMBAR_ATENCION_FONDO)


def test_registrar_vigencia_y_fuente_desde_la_ventana_principal(qtbot, raiz_temporal, monkeypatch) -> None:
    """Modulo 11 (Vigencias y fuentes): registrar desde la ventana deja
    consultable tanto la Fuente como la Vigencia del documento, y el
    panel deja de mostrar "Sin registrar."."""
    raiz, origenes_dir = raiz_temporal
    ctx = bootstrap(raiz=str(raiz))
    proyecto = ctx.projects.crear_proyecto(nombre="GUI-vigencia", perfil_id="perfil_generico")

    archivo = origenes_dir / "contrato.docx"
    d = docx.Document()
    d.add_heading("Contrato de prueba", level=1)
    d.add_paragraph("Contenido de prueba para la GUI.")
    d.save(str(archivo))

    ventana = MainWindow(ctx)
    qtbot.addWidget(ventana)
    ventana._on_abrir_proyecto()

    documento = ventana.pctx.documents.registrar_documento(proyecto_id=proyecto.id, ruta_origen=str(archivo))
    ventana.pctx.documents.clasificar_documento(documento.id)
    extraido = ventana.pctx.documents.extraer_documento(documento.id)

    _ir_a_revision(ventana)
    ventana.tabla_documentos.selectRow(0)
    assert ventana.documento_actual.id == extraido.id
    assert ventana.etiqueta_vigencia_fuente.text() == "Sin registrar."
    assert ventana.boton_registrar_vigencia_fuente.isEnabled()

    def _dialogo_vigencia_prellenado(self, parent=None, fuente=None, vigencia=None, documentos_disponibles=None):
        RegistrarVigenciaDialog.__base_init__(self, parent, fuente, vigencia, documentos_disponibles)
        self.campo_sistema_origen.setText("portal de normatividad interna")
        self.campo_autoridad_fuente.setText("Direccion Juridica")
        self.campo_evidencia_consulta.setText("Confirmado el 2026-07-28 con el area emisora.")
        self.campo_vigente_desde.setText("2026-01-01")
        self.campo_sin_vencimiento.setChecked(True)
        self.campo_autoridad_vigencia.setText("Direccion Juridica")

    RegistrarVigenciaDialog.__base_init__ = RegistrarVigenciaDialog.__init__
    monkeypatch.setattr(RegistrarVigenciaDialog, "__init__", _dialogo_vigencia_prellenado)
    monkeypatch.setattr(RegistrarVigenciaDialog, "exec", lambda self: 1)

    ventana._on_registrar_vigencia_fuente()

    fuente = ventana.pctx.validity.obtener_fuente(extraido.id)
    vigencia = ventana.pctx.validity.obtener_vigencia(extraido.id)
    assert fuente is not None
    assert fuente.sistema_origen == "portal de normatividad interna"
    assert vigencia is not None
    assert vigencia.esta_completa() is True
    assert "Direccion Juridica" in ventana.etiqueta_vigencia_fuente.text()
    assert "sin vencimiento" in ventana.etiqueta_vigencia_fuente.text()


def test_ventana_arranca_sin_proyectos_sin_fallar(qtbot) -> None:
    """Ventana sin ningun proyecto creado -- no debe fallar al
    construirse, solo mostrar el combo vacio."""
    raiz = Path("/tmp/rag_content_studio_smoke_test_gui_vacio")
    if raiz.exists():
        shutil.rmtree(raiz)
    raiz.mkdir(parents=True)
    try:
        ctx = bootstrap(raiz=str(raiz))
        ventana = MainWindow(ctx)
        qtbot.addWidget(ventana)
        assert ventana.combo_proyectos.count() == 0
        assert ventana.pctx is None
    finally:
        shutil.rmtree(raiz, ignore_errors=True)
