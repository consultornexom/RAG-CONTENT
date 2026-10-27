"""Prueba de humo de la Fase 11 (Interfaz completa, Documento Maestro
seccion 4.11): recorre el menu funcional nuevo de punta a punta --
bootstrap()/abrir_proyecto() reales, sin fakes, mismo patron que
tests/smoke/test_gui_revision_incidencias_extremo_a_extremo.py (Modulo
8, sin cambios). Los dialogos modales se parchan igual que alli, para
no bloquear la prueba en modo headless.

`ProyectosPanel`/`ConfiguracionPanel` (respaldos, Fase 10) usan
`keyring` -- este sandbox de pruebas no tiene sesion de escritorio real
(`keyring.errors.NoKeyringError`), asi que se sustituye el backend por
uno en memoria, mismo criterio que
tests/integration/test_respaldo_extremo_a_extremo.py (movido ahi en la
cuarta ronda de correcciones, 2026-08-01: no depende de PySide6/qtbot,
ver docstring de tests/integration/__init__.py)."""
from __future__ import annotations

import shutil
from pathlib import Path
from typing import Optional

import docx
import keyring
import keyring.backend
import keyring.errors
import pytest

pytest.importorskip("PySide6")

from PySide6.QtWidgets import QFileDialog, QMessageBox  # noqa: E402

from app.bootstrap import bootstrap  # noqa: E402
from domain.documents.entities import EstadoDocumento  # noqa: E402
from domain.incidents.entities import EstadoIncidencia, EtapaIncidencia, SeveridadIncidencia  # noqa: E402
from ui.dialogs import (  # noqa: E402
    AprobadorDialog,
    CrearPerfilDialog,
    CrearProyectoDialog,
    DecidirIncidenciaDialog,
    RegistrarVigenciaDialog,
    RevisarDocumentoDialog,
    TextoSimpleDialog,
)
from ui.main_window import MainWindow  # noqa: E402

RAIZ_PRUEBA = Path("/tmp/rag_content_studio_smoke_test_menu_completo")


class _BackendEnMemoria(keyring.backend.KeyringBackend):
    priority = 1  # type: ignore[assignment]

    def __init__(self) -> None:
        super().__init__()
        self._datos: dict[tuple[str, str], str] = {}

    def set_password(self, servicename: str, username: str, password: str) -> None:
        self._datos[(servicename, username)] = password

    def get_password(self, servicename: str, username: str) -> Optional[str]:
        return self._datos.get((servicename, username))

    def delete_password(self, servicename: str, username: str) -> None:
        if (servicename, username) not in self._datos:
            raise keyring.errors.PasswordDeleteError("no existe")
        del self._datos[(servicename, username)]


@pytest.fixture()
def backend_en_memoria():
    backend_original = keyring.get_keyring()
    keyring.set_keyring(_BackendEnMemoria())
    yield
    keyring.set_keyring(backend_original)


@pytest.fixture()
def raiz_temporal():
    if RAIZ_PRUEBA.exists():
        shutil.rmtree(RAIZ_PRUEBA)
    RAIZ_PRUEBA.mkdir(parents=True)
    perfiles_dir = RAIZ_PRUEBA / "perfiles" / "perfil_generico"
    perfiles_dir.mkdir(parents=True)
    shutil.copy(
        Path(__file__).resolve().parents[2] / "perfiles" / "perfil_generico" / "0.1.yaml",
        perfiles_dir / "0.1.yaml",
    )
    origenes_dir = RAIZ_PRUEBA / "_origenes_de_prueba"
    origenes_dir.mkdir()
    yield RAIZ_PRUEBA, origenes_dir
    shutil.rmtree(RAIZ_PRUEBA, ignore_errors=True)


def _indice_pagina(ventana: MainWindow, etiqueta: str) -> int:
    for i in range(ventana.lista_navegacion.count()):
        if ventana.lista_navegacion.item(i).text() == etiqueta:
            return i
    raise AssertionError(f"No existe la pagina '{etiqueta}' en el menu")


def _ir_a(ventana: MainWindow, etiqueta: str) -> None:
    ventana.lista_navegacion.setCurrentRow(_indice_pagina(ventana, etiqueta))


def test_recorrido_completo_del_documento_por_el_menu_nuevo(qtbot, raiz_temporal, monkeypatch, backend_en_memoria) -> None:
    raiz, origenes_dir = raiz_temporal
    ctx = bootstrap(raiz=str(raiz))

    ventana = MainWindow(ctx)
    qtbot.addWidget(ventana)
    assert ventana.paginas.count() == 13

    # -- Proyectos: crear y abrir automaticamente.
    _ir_a(ventana, "Proyectos")

    def _dialogo_proyecto_prellenado(self, parent=None, perfiles=None):
        CrearProyectoDialog.__base_init__(self, parent, perfiles)
        self.campo_nombre.setText("Fase 11")
        indice = self.campo_perfil.findData("perfil_generico")
        self.campo_perfil.setCurrentIndex(indice)

    CrearProyectoDialog.__base_init__ = CrearProyectoDialog.__init__
    monkeypatch.setattr(CrearProyectoDialog, "__init__", _dialogo_proyecto_prellenado)
    monkeypatch.setattr(CrearProyectoDialog, "exec", lambda self: 1)

    ventana.panel_proyectos._on_crear_proyecto()
    assert ventana.pctx is not None
    assert ventana.proyecto_actual is not None
    assert ventana.proyecto_actual.nombre == "Fase 11"
    assert ventana.panel_proyectos.tabla.rowCount() == 1

    # -- Inventario: agregar un documento real via el selector de archivos (parchado).
    _ir_a(ventana, "Inventario")
    archivo = origenes_dir / "politica.docx"
    d = docx.Document()
    d.add_heading("Politica de prueba", level=1)
    d.add_paragraph("Contenido real para la Fase 11.")
    d.save(str(archivo))
    monkeypatch.setattr(QFileDialog, "getOpenFileNames", staticmethod(lambda *a, **k: ([str(archivo)], "")))
    ventana.panel_inventario._on_agregar_documentos()
    assert ventana.panel_inventario.tabla.rowCount() == 1
    documento = ventana.pctx.documents.listar_documentos()[0]

    # -- Procesamiento: clasificar, extraer sincrono, y encolar+procesar en segundo plano.
    _ir_a(ventana, "Procesamiento")
    ventana.panel_procesamiento.tabla_documentos.selectRow(0)
    ventana.panel_procesamiento._on_clasificar()
    ventana.panel_procesamiento._on_extraer_sincrono()
    documento = ventana.pctx.documents.obtener_documento(documento.id)
    assert documento.estado == EstadoDocumento.EXTRAIDO

    # -- Revision (Modulo 8, sin cambios): aprobar para poder normalizar despues.
    # Bug real corregido 2026-07-30: antes de esta fecha, entrar a esta
    # pestana no refrescaba `tabla_documentos` -- todavia mostraria 0
    # filas (la foto de cuando se abrio el proyecto, antes de registrar
    # y procesar el documento arriba), y esta prueba necesitaba forzar
    # un `ventana._cargar_documentos()` manual para poder seleccionar la
    # fila. Ya no hace falta: `_ir_a()` (que simplemente cambia de
    # pestana, como haria un operador real) debe dejar la tabla al dia
    # por si sola.
    _ir_a(ventana, "Revision")
    assert ventana.tabla_documentos.rowCount() == 1
    ventana.tabla_documentos.selectRow(0)
    assert ventana.documento_actual.id == documento.id

    def _dialogo_revision_prellenado(self, parent=None):
        RevisarDocumentoDialog.__base_init__(self, parent)
        self.campo_revisor.setText("ana@nexom.mx")

    RevisarDocumentoDialog.__base_init__ = RevisarDocumentoDialog.__init__
    monkeypatch.setattr(RevisarDocumentoDialog, "__init__", _dialogo_revision_prellenado)
    monkeypatch.setattr(RevisarDocumentoDialog, "exec", lambda self: 1)
    ventana._on_aprobar_revision()
    assert ventana.documento_actual.estado == EstadoDocumento.REVISADO

    # -- Vigencia y fuente (Modulo 11, sigue viviendo en "Revision"):
    # liberar_documento() exige una Fuente registrada
    # (application.governance.verificar_documento_liberable), asi que
    # esto tiene que pasar antes de llegar a "Markdown" mas abajo.
    def _dialogo_vigencia_prellenado(self, parent=None, fuente=None, vigencia=None, documentos_disponibles=None):
        RegistrarVigenciaDialog.__base_init__(self, parent, fuente, vigencia, documentos_disponibles)
        self.campo_sistema_origen.setText("portal de normatividad interna")
        self.campo_autoridad_fuente.setText("Direccion Juridica")
        self.campo_evidencia_consulta.setText("Confirmado el 2026-07-28.")
        self.campo_sin_vencimiento.setChecked(True)
        self.campo_autoridad_vigencia.setText("Direccion Juridica")

    RegistrarVigenciaDialog.__base_init__ = RegistrarVigenciaDialog.__init__
    monkeypatch.setattr(RegistrarVigenciaDialog, "__init__", _dialogo_vigencia_prellenado)
    monkeypatch.setattr(RegistrarVigenciaDialog, "exec", lambda self: 1)
    ventana._on_registrar_vigencia_fuente()
    assert ventana.pctx.validity.obtener_fuente(documento.id) is not None

    # -- Normalizacion: primera vez que se dispara desde la interfaz.
    _ir_a(ventana, "Normalizacion")
    ventana.panel_normalizacion.tabla.selectRow(0)
    assert ventana.panel_normalizacion.boton_normalizar.isEnabled()
    ventana.panel_normalizacion._on_normalizar()
    documento = ventana.pctx.documents.obtener_documento(documento.id)
    assert documento.estado == EstadoDocumento.NORMALIZADO

    # -- Markdown: aprobar, generar, validar, liberar y generar manifiesto.
    _ir_a(ventana, "Markdown")
    ventana.panel_markdown.tabla.selectRow(0)

    def _dialogo_aprobador_prellenado(self, parent=None, titulo="Aprobar"):
        AprobadorDialog.__base_init__(self, parent, titulo)
        self.campo_aprobador.setText("ana@nexom.mx")

    AprobadorDialog.__base_init__ = AprobadorDialog.__init__
    monkeypatch.setattr(AprobadorDialog, "__init__", _dialogo_aprobador_prellenado)
    monkeypatch.setattr(AprobadorDialog, "exec", lambda self: 1)
    ventana.panel_markdown._on_aprobar()
    ventana.panel_markdown._on_generar()
    ventana.panel_markdown._on_validar()
    ventana.panel_markdown._on_liberar()
    documento = ventana.pctx.documents.obtener_documento(documento.id)
    assert documento.estado == EstadoDocumento.LIBERADO

    monkeypatch.setattr("ui.panels.markdown_panel.QDesktopServices.openUrl", lambda *a, **k: None)
    # QMessageBox.information() de exito es un dialogo modal real -- sin
    # parchar, se queda esperando un clic que nunca llega en un entorno
    # headless (offscreen).
    monkeypatch.setattr("ui.panels.markdown_panel.QMessageBox.information", lambda *a, **k: None)
    ventana.panel_markdown._on_preparar_entrega_rag_studio()

    # -- Vigencias: la tabla del proyecto ya muestra el documento (no es critico, no bloquea nada).
    _ir_a(ventana, "Vigencias")
    assert ventana.panel_vigencias.tabla.rowCount() == 1

    # -- Lotes: crear el lote desde el panel nuevo (hallazgo real de la
    # cuarta auditoria general, 2026-08-01: Modo Avanzado no tenia forma
    # de crear/cerrar un lote fuera del asistente por etapas -- ver
    # docstring de ui.panels.lotes_panel), agregar el documento via la
    # fachada (esta pantalla no expone agregar/quitar documentos, solo
    # crear/cerrar) y cerrarlo tambien desde el panel.
    _ir_a(ventana, "Lotes")

    def _dialogo_nuevo_lote_prellenado(self, parent=None, titulo="", etiqueta="", una_linea=False):
        TextoSimpleDialog.__base_init__(self, parent, titulo, etiqueta, una_linea)
        self.campo.setText("Lote de prueba")

    TextoSimpleDialog.__base_init__ = TextoSimpleDialog.__init__
    monkeypatch.setattr(TextoSimpleDialog, "__init__", _dialogo_nuevo_lote_prellenado)
    monkeypatch.setattr(TextoSimpleDialog, "exec", lambda self: 1)
    ventana.panel_lotes._on_nuevo()
    assert ventana.panel_lotes.tabla.rowCount() == 1
    lote = ventana.pctx.batches.listar_lotes(ventana.proyecto_actual.id)[0]
    assert lote.estado.value == "abierto"

    ventana.pctx.batches.agregar_documento(lote.id, documento.id)
    ventana.panel_lotes._cargar_tabla()
    assert ventana.panel_lotes.tabla.item(0, 2).text() == "1"  # columna "Documentos"

    monkeypatch.setattr(
        "ui.panels.lotes_panel.QMessageBox.question", lambda *a, **k: QMessageBox.StandardButton.Yes
    )
    ventana.panel_lotes.tabla.selectRow(0)
    ventana.panel_lotes._on_cerrar()
    lote = ventana.pctx.batches.obtener_lote(lote.id)
    assert lote.estado.value == "cerrado"

    # -- Reportes: generar los seis informes + acta de cierre del lote de arriba.
    _ir_a(ventana, "Reportes")
    assert ventana.panel_reportes.combo_lotes.count() == 1
    ventana.panel_reportes._on_generar_todos()
    assert ventana.panel_reportes.lista_archivos.count() == 6
    ventana.panel_reportes.combo_lotes.setCurrentIndex(0)
    ventana.panel_reportes._on_generar_acta()
    assert ventana.panel_reportes.lista_archivos.count() == 7

    # -- Auditoria: ya hay eventos reales registrados por todo lo anterior.
    _ir_a(ventana, "Auditoria")
    assert ventana.panel_auditoria.tabla.rowCount() > 0

    # -- Configuracion: crear un perfil nuevo y un respaldo cifrado real del proyecto.
    _ir_a(ventana, "Configuracion")
    total_perfiles_antes = ventana.panel_configuracion.lista_perfiles.count()

    def _dialogo_perfil_prellenado(self, parent=None):
        CrearPerfilDialog.__base_init__(self, parent)
        self.campo_perfil_id.setText("perfil_fase_11")

    CrearPerfilDialog.__base_init__ = CrearPerfilDialog.__init__
    monkeypatch.setattr(CrearPerfilDialog, "__init__", _dialogo_perfil_prellenado)
    monkeypatch.setattr(CrearPerfilDialog, "exec", lambda self: 1)
    ventana.panel_configuracion._on_crear_perfil()
    assert ventana.panel_configuracion.lista_perfiles.count() == total_perfiles_antes + 1

    ventana.panel_configuracion._on_crear_respaldo()
    assert "Respaldo creado" in ventana.panel_configuracion.etiqueta_respaldo.text()


def test_pagina_incidencias_ciclo_completo(qtbot, raiz_temporal, monkeypatch) -> None:
    raiz, origenes_dir = raiz_temporal
    ctx = bootstrap(raiz=str(raiz))
    proyecto = ctx.projects.crear_proyecto(nombre="Fase 11 incidencias", perfil_id="perfil_generico")

    archivo = origenes_dir / "memo.txt"
    archivo.write_text("Memo real de prueba.", encoding="utf-8")

    ventana = MainWindow(ctx)
    qtbot.addWidget(ventana)
    ventana.combo_proyectos.setCurrentIndex(ventana.combo_proyectos.findData(proyecto.id))
    ventana._on_abrir_proyecto()

    documento = ventana.pctx.documents.registrar_documento(proyecto_id=proyecto.id, ruta_origen=str(archivo))
    incidencia = ventana.pctx.incidents.registrar_incidencia(
        proyecto_id=proyecto.id,
        documento_id=documento.id,
        etapa=EtapaIncidencia.INVENTARIO,
        tipo="dato_faltante",
        severidad=SeveridadIncidencia.MEDIA,
        evidencia="falta el area propietaria",
        impacto="no se puede clasificar con certeza",
    )

    _ir_a(ventana, "Incidencias")
    assert ventana.panel_incidencias.tabla.rowCount() == 1
    ventana.panel_incidencias.tabla.selectRow(0)

    monkeypatch.setattr(TextoSimpleDialog, "exec", lambda self: 1)

    def _asignar_prellenado(self, parent=None, titulo="", etiqueta="", una_linea=False):
        TextoSimpleDialog.__base_init__(self, parent, titulo, etiqueta, una_linea)
        self.campo.setText("ana@nexom.mx")

    TextoSimpleDialog.__base_init__ = TextoSimpleDialog.__init__
    monkeypatch.setattr(TextoSimpleDialog, "__init__", _asignar_prellenado)
    ventana.panel_incidencias._on_asignar()
    assert ventana.pctx.incidents.obtener_incidencia(incidencia.id).responsable == "ana@nexom.mx"

    def _responder_prellenado(self, parent=None, titulo="", etiqueta="", una_linea=False):
        TextoSimpleDialog.__base_init__(self, parent, titulo, etiqueta, una_linea)
        self.campo.setPlainText("Se agrego el area propietaria.")

    monkeypatch.setattr(TextoSimpleDialog, "__init__", _responder_prellenado)
    ventana.panel_incidencias._on_responder()

    def _dialogo_decidir_prellenado(self, parent=None):
        DecidirIncidenciaDialog.__base_init__(self, parent)
        self.campo_decision.setPlainText("Corregido, se acepta la respuesta.")
        self.campo_resuelta.setChecked(True)

    DecidirIncidenciaDialog.__base_init__ = DecidirIncidenciaDialog.__init__
    monkeypatch.setattr(DecidirIncidenciaDialog, "__init__", _dialogo_decidir_prellenado)
    monkeypatch.setattr(DecidirIncidenciaDialog, "exec", lambda self: 1)
    ventana.panel_incidencias._on_decidir()
    assert ventana.pctx.incidents.obtener_incidencia(incidencia.id).estado == EstadoIncidencia.RESUELTA

    def _cerrar_prellenado(self, parent=None, titulo="", etiqueta="", una_linea=False):
        TextoSimpleDialog.__base_init__(self, parent, titulo, etiqueta, una_linea)
        self.campo.setPlainText("Incidencia cerrada tras corregirse.")

    monkeypatch.setattr(TextoSimpleDialog, "__init__", _cerrar_prellenado)
    ventana.panel_incidencias._on_cerrar()
    assert ventana.pctx.incidents.obtener_incidencia(incidencia.id).esta_abierta() is False


def test_pagina_incidencias_avisa_si_se_confirma_con_el_campo_obligatorio_vacio(
    qtbot, raiz_temporal, monkeypatch
) -> None:
    """Bug real encontrado en la segunda auditoria 2026-07-31:
    IncidenciasPanel era el unico panel de toda la interfaz que, si el
    operador confirmaba el dialogo (Asignar/Responder/Decidir/Cerrar)
    con el campo obligatorio vacio, no avisaba nada -- simplemente no
    pasaba nada, sin ninguna señal de por que. Corregido con
    QMessageBox.warning("Campo requerido", ...), igual que el resto de
    la interfaz. Se prueban los cuatro botones."""
    raiz, origenes_dir = raiz_temporal
    ctx = bootstrap(raiz=str(raiz))
    proyecto = ctx.projects.crear_proyecto(nombre="Fase 11 incidencias vacias", perfil_id="perfil_generico")

    archivo = origenes_dir / "memo.txt"
    archivo.write_text("Memo real de prueba.", encoding="utf-8")

    ventana = MainWindow(ctx)
    qtbot.addWidget(ventana)
    ventana.combo_proyectos.setCurrentIndex(ventana.combo_proyectos.findData(proyecto.id))
    ventana._on_abrir_proyecto()

    documento = ventana.pctx.documents.registrar_documento(proyecto_id=proyecto.id, ruta_origen=str(archivo))
    incidencia = ventana.pctx.incidents.registrar_incidencia(
        proyecto_id=proyecto.id,
        documento_id=documento.id,
        etapa=EtapaIncidencia.INVENTARIO,
        tipo="dato_faltante",
        severidad=SeveridadIncidencia.MEDIA,
        evidencia="falta el area propietaria",
        impacto="no se puede clasificar con certeza",
    )

    _ir_a(ventana, "Incidencias")
    ventana.panel_incidencias.tabla.selectRow(0)

    llamadas_warning: list[tuple] = []
    monkeypatch.setattr(QMessageBox, "warning", lambda *a, **k: llamadas_warning.append(a))

    # Los cuatro dialogos se confirman (exec() -> 1) con su campo
    # obligatorio vacio -- ninguno modifica su __init__, se dejan con el
    # valor por defecto (vacio).
    monkeypatch.setattr(TextoSimpleDialog, "exec", lambda self: 1)
    monkeypatch.setattr(DecidirIncidenciaDialog, "exec", lambda self: 1)

    ventana.panel_incidencias._on_asignar()
    assert ventana.pctx.incidents.obtener_incidencia(incidencia.id).responsable is None
    ventana.panel_incidencias._on_responder()
    ventana.panel_incidencias._on_decidir()
    assert ventana.pctx.incidents.obtener_incidencia(incidencia.id).estado == EstadoIncidencia.ABIERTA
    ventana.panel_incidencias._on_cerrar()
    assert ventana.pctx.incidents.obtener_incidencia(incidencia.id).esta_abierta() is True

    assert len(llamadas_warning) == 4
    assert all(args[1] == "Campo requerido" for args in llamadas_warning)


def test_pagina_pruebas_de_aceptacion_corre_contra_un_corpus_real(qtbot, raiz_temporal, monkeypatch) -> None:
    raiz, origenes_dir = raiz_temporal
    ctx = bootstrap(raiz=str(raiz))
    proyecto = ctx.projects.crear_proyecto(nombre="Fase 11 pruebas", perfil_id="perfil_generico")

    corpus_dir = raiz / "_corpus_prueba"
    corpus_dir.mkdir()
    (corpus_dir / "memo.txt").write_text("Contenido real del corpus de prueba.", encoding="utf-8")

    ventana = MainWindow(ctx)
    qtbot.addWidget(ventana)
    ventana.combo_proyectos.setCurrentIndex(ventana.combo_proyectos.findData(proyecto.id))
    ventana._on_abrir_proyecto()

    _ir_a(ventana, "Pruebas de aceptacion")
    monkeypatch.setattr(QFileDialog, "getExistingDirectory", staticmethod(lambda *a, **k: str(corpus_dir)))
    ventana.panel_pruebas._on_elegir_corpus()
    assert ventana.panel_pruebas.boton_ejecutar.isEnabled()

    ventana.panel_pruebas._on_ejecutar()
    assert "Resultado:" in ventana.panel_pruebas.etiqueta_resultado.text()
    assert ventana.panel_pruebas.lista_dominios.count() == 7
    ruta_informe = Path(proyecto.ruta_carpeta) / "10_reportes"
    assert any(ruta_informe.glob("informe_aceptacion_*.md"))


def test_navegar_entre_paginas_no_pierde_el_estado_de_revision(qtbot, raiz_temporal) -> None:
    """El cambio de pagina refresca los paneles con contexto (Fase 11)
    sin romper la pagina de Revision, que sigue su propio ciclo de vida
    (Modulo 8) sin cambios."""
    raiz, _origenes_dir = raiz_temporal
    ctx = bootstrap(raiz=str(raiz))
    ctx.projects.crear_proyecto(nombre="Fase 11 navegacion", perfil_id="perfil_generico")

    ventana = MainWindow(ctx)
    qtbot.addWidget(ventana)
    ventana._on_abrir_proyecto()

    for etiqueta in ["Inventario", "Procesamiento", "Auditoria", "Revision", "Lotes", "Configuracion"]:
        _ir_a(ventana, etiqueta)
        assert ventana.paginas.currentIndex() == _indice_pagina(ventana, etiqueta)

    assert ventana.pctx is not None


def test_cambiar_de_modo_refresca_el_asistente_y_el_modo_avanzado(qtbot, raiz_temporal, monkeypatch) -> None:
    """Bug real, auditoria 2026-08-24: MainWindow._on_cambiar_modo solo
    refrescaba "Inicio" al cambiar de modo -- si el operador ya estaba
    parado en el asistente o en la misma pagina del modo avanzado antes
    de cambiar de modo, y volvia a ese mismo lugar despues, no se
    disparaba ningun refresco: ni AsistenteWidget._mostrar_etapa ni
    MainWindow._on_cambiar_pagina dependen de que la pestana se VUELVA a
    mostrar, solo de que el indice/renglon efectivamente CAMBIE. Se
    confirma aqui, con espias sobre los metodos de refresco reales
    (no dobles), que entrar a "Asistente" y a "Modo avanzado" dispara su
    refresco correspondiente incluso sin cambiar de etapa/renglon."""
    raiz, _origenes_dir = raiz_temporal
    ctx = bootstrap(raiz=str(raiz))
    ctx.projects.crear_proyecto(nombre="Fase 11 refresco de modo", perfil_id="perfil_generico")

    ventana = MainWindow(ctx)
    qtbot.addWidget(ventana)
    ventana._on_abrir_proyecto()  # deja la ventana en "Asistente" (vista predeterminada al abrir).
    assert ventana.radio_asistente.isChecked()

    llamadas_asistente: list[object] = []
    establecer_proyecto_original = ventana.asistente.establecer_proyecto

    def _espia_asistente(pctx, proyecto):
        llamadas_asistente.append(proyecto)
        return establecer_proyecto_original(pctx, proyecto)

    monkeypatch.setattr(ventana.asistente, "establecer_proyecto", _espia_asistente)

    llamadas_pagina: list[int] = []
    on_cambiar_pagina_original = ventana._on_cambiar_pagina

    def _espia_pagina(indice):
        llamadas_pagina.append(indice)
        return on_cambiar_pagina_original(indice)

    monkeypatch.setattr(ventana, "_on_cambiar_pagina", _espia_pagina)

    # "Modo avanzado" ya queda parado en el primer renglon de
    # lista_navegacion (sin cambiarlo) -- antes del fix, esto NO
    # disparaba ningun refresco de pagina.
    ventana.radio_modo_avanzado.click()
    assert len(llamadas_pagina) >= 1

    # Volver a "Asistente" (misma etapa autoguardada, sin cambiarla) --
    # antes del fix, esto tampoco refrescaba el asistente.
    ventana.radio_asistente.click()
    assert len(llamadas_asistente) >= 1


def test_boton_abrir_carpeta_proyectos_muestra_error_amigable_si_no_hay_permiso(
    qtbot, raiz_temporal, monkeypatch
) -> None:
    """Bug real, auditoria 2026-08-24:
    ConfiguracionPanel._on_abrir_carpeta_proyectos no envolvia su
    `mkdir()` en try/except, a diferencia de todos los demas puntos de
    escritura en disco de ese mismo panel (_on_crear_respaldo,
    _on_restaurar_respaldo) -- un PermissionError/OSError real (carpeta
    raiz de datos sin permiso de escritura) se propagaba crudo desde un
    slot de Qt en vez de mostrarse con un dialogo amigable, como el
    resto de la interfaz."""
    raiz, _origenes_dir = raiz_temporal
    ctx = bootstrap(raiz=str(raiz))
    ventana = MainWindow(ctx)
    qtbot.addWidget(ventana)
    _ir_a(ventana, "Configuracion")

    mensajes: list[str] = []
    monkeypatch.setattr(
        "ui.panels.configuracion_panel.QMessageBox.critical",
        lambda self, titulo, texto: mensajes.append(texto),
    )
    monkeypatch.setattr(
        "ui.panels.configuracion_panel.Path.mkdir",
        lambda self, *a, **k: (_ for _ in ()).throw(PermissionError("sin permiso de escritura (simulado)")),
    )

    ventana.panel_configuracion._on_abrir_carpeta_proyectos()

    assert len(mensajes) == 1
