"""DocLegacyExtractor y XlsLegacyExtractor (conversion DOC/XLS legado)
contra un LibreOfficeConverter real -- mismo criterio que
test_ocr_extractors.py: lo que estas pruebas necesitan verificar es la
integracion real (conversion + delegacion en DocxExtractor/
XlsxExtractor), no un doble de LibreOffice.
"""
from __future__ import annotations

import shutil
import subprocess
import tempfile
from pathlib import Path

import docx
import openpyxl
import pytest

from domain.documents.entities import FormatoDocumento, TipoProcesamiento
from infrastructure.extraction.doc_legacy_extractor import DocLegacyExtractor
from infrastructure.extraction.libreoffice_converter import LibreOfficeConverter
from infrastructure.extraction.xls_legacy_extractor import XlsLegacyExtractor

RAIZ_PRUEBA = Path("/tmp/rag_content_studio_legacy_extractores_test")


@pytest.fixture()
def raiz():
    if RAIZ_PRUEBA.exists():
        shutil.rmtree(RAIZ_PRUEBA)
    RAIZ_PRUEBA.mkdir(parents=True)
    yield RAIZ_PRUEBA
    shutil.rmtree(RAIZ_PRUEBA, ignore_errors=True)


@pytest.fixture()
def converter() -> LibreOfficeConverter:
    return LibreOfficeConverter()


def _convertir_con_soffice(ruta_entrada: Path, filtro_destino: str, directorio_salida: Path) -> None:
    perfil = tempfile.mkdtemp(prefix="rag_content_studio_test_lo_profile_")
    try:
        subprocess.run(
            [
                "soffice",
                "--headless",
                "--norestore",
                f"-env:UserInstallation=file://{perfil}",
                "--convert-to",
                filtro_destino,
                "--outdir",
                str(directorio_salida),
                str(ruta_entrada),
            ],
            capture_output=True,
            text=True,
            timeout=120,
            check=True,
        )
    finally:
        shutil.rmtree(perfil, ignore_errors=True)


def _doc_de_prueba(raiz: Path) -> Path:
    d = docx.Document()
    d.add_heading("Contrato legado", level=1)
    d.add_paragraph("Texto de prueba en formato DOC clasico.")
    tabla = d.add_table(rows=2, cols=2)
    tabla.rows[0].cells[0].text = "Clausula"
    tabla.rows[0].cells[1].text = "Descripcion"
    tabla.rows[1].cells[0].text = "1"
    tabla.rows[1].cells[1].text = "Vigencia"
    docx_path = raiz / "contrato.docx"
    d.save(str(docx_path))
    _convertir_con_soffice(docx_path, "doc", raiz)
    return raiz / "contrato.doc"


def _xls_de_prueba(raiz: Path) -> Path:
    libro = openpyxl.Workbook()
    hoja = libro.active
    hoja.title = "Datos"
    hoja.append(["nombre", "area"])
    hoja.append(["Ana", "RH"])
    xlsx_path = raiz / "datos.xlsx"
    libro.save(str(xlsx_path))
    _convertir_con_soffice(xlsx_path, "xls", raiz)
    return raiz / "datos.xls"


def test_doc_legacy_extractor_supports_solo_doc(converter: LibreOfficeConverter) -> None:
    extractor = DocLegacyExtractor(converter)
    assert extractor.supports(FormatoDocumento.DOC, TipoProcesamiento.EXTRACCION_DIRECTA) is True
    assert extractor.supports(FormatoDocumento.DOCX, TipoProcesamiento.EXTRACCION_DIRECTA) is False


def test_doc_legacy_extractor_convierte_y_extrae_contenido_real(raiz: Path, converter: LibreOfficeConverter) -> None:
    ruta_doc = _doc_de_prueba(raiz)

    resultado = DocLegacyExtractor(converter).extract(
        str(ruta_doc), FormatoDocumento.DOC, TipoProcesamiento.EXTRACCION_DIRECTA
    )
    # No se exige la tabla como tabla Markdown real: el roundtrip
    # DOCX->DOC->DOCX vía LibreOffice puede aplanar una tabla simple a
    # texto corrido (artefacto real de la conversion, no de este
    # extractor -- DocxExtractor si produce tabla Markdown cuando el
    # DOCX de origen la conserva como tabla, ver
    # test_document_extractors.py). Lo que si es responsabilidad de
    # este extractor es que el texto sobreviva integro al roundtrip.
    assert "# Contrato legado" in resultado.markdown
    assert "Texto de prueba en formato DOC clasico." in resultado.markdown
    assert "Clausula" in resultado.markdown and "Vigencia" in resultado.markdown
    assert resultado.metodo == "doc_legado_via_docx"


def test_doc_legacy_extractor_no_deja_directorios_temporales(raiz: Path, converter: LibreOfficeConverter) -> None:
    ruta_doc = _doc_de_prueba(raiz)
    tmp_antes = set(Path(tempfile.gettempdir()).glob("rag_content_studio_legacy_*"))

    DocLegacyExtractor(converter).extract(str(ruta_doc), FormatoDocumento.DOC, TipoProcesamiento.EXTRACCION_DIRECTA)

    tmp_despues = set(Path(tempfile.gettempdir()).glob("rag_content_studio_legacy_*"))
    assert tmp_despues == tmp_antes


def test_xls_legacy_extractor_supports_solo_xls(converter: LibreOfficeConverter) -> None:
    extractor = XlsLegacyExtractor(converter)
    assert extractor.supports(FormatoDocumento.XLS, TipoProcesamiento.EXTRACCION_DIRECTA) is True
    assert extractor.supports(FormatoDocumento.XLSX, TipoProcesamiento.EXTRACCION_DIRECTA) is False


def test_xls_legacy_extractor_convierte_y_extrae_contenido_real(raiz: Path, converter: LibreOfficeConverter) -> None:
    ruta_xls = _xls_de_prueba(raiz)

    resultado = XlsLegacyExtractor(converter).extract(
        str(ruta_xls), FormatoDocumento.XLS, TipoProcesamiento.EXTRACCION_DIRECTA
    )
    assert "| nombre | area |" in resultado.markdown
    assert "| Ana | RH |" in resultado.markdown
    assert resultado.metodo == "xls_legado_via_xlsx"


def test_doc_legacy_extractor_archivo_ilegible_propaga_error_de_conversion(
    raiz: Path, converter: LibreOfficeConverter
) -> None:
    from domain.documents.exceptions import LegacyConversionFailedError

    # LibreOffice es muy permisivo con la deteccion de formato -- texto
    # plano renombrado a .doc se convierte igual (autodetecta "texto"),
    # asi que hace falta imitar bytes de cabecera OLE2 reales seguidos
    # de basura para forzar un fallo real de carga, no un simple .doc
    # con contenido atipico.
    ruta = raiz / "no_es_doc.doc"
    cabecera_ole2 = bytes([0xD0, 0xCF, 0x11, 0xE0, 0xA1, 0xB1, 0x1A, 0xE1])
    ruta.write_bytes(cabecera_ole2 + b"\x00" * 500 + bytes(range(256)) * 3)

    with pytest.raises(LegacyConversionFailedError):
        DocLegacyExtractor(converter).extract(str(ruta), FormatoDocumento.DOC, TipoProcesamiento.EXTRACCION_DIRECTA)
