from __future__ import annotations

import json
import shutil
from pathlib import Path

import docx
import openpyxl
import pytest
from reportlab.pdfgen import canvas

from domain.documents.entities import FormatoDocumento, TipoProcesamiento
from domain.documents.exceptions import ExtractionFailedError, UnsupportedFormatError
from infrastructure.extraction.composite_document_extractor import CompositeDocumentExtractor
from infrastructure.extraction.csv_extractor import CsvExtractor
from infrastructure.extraction.docx_extractor import DocxExtractor
from infrastructure.extraction.json_extractor import JsonExtractor
from infrastructure.extraction.pdf_extractor import PdfExtractor
from infrastructure.extraction.txt_markdown_extractor import TxtMarkdownExtractor
from infrastructure.extraction.xlsx_extractor import XlsxExtractor

RAIZ_PRUEBA = Path("/tmp/rag_content_studio_extractores_test")


@pytest.fixture()
def raiz():
    if RAIZ_PRUEBA.exists():
        shutil.rmtree(RAIZ_PRUEBA)
    RAIZ_PRUEBA.mkdir(parents=True)
    yield RAIZ_PRUEBA
    shutil.rmtree(RAIZ_PRUEBA, ignore_errors=True)


def test_txt_extractor_lee_texto_literal(raiz: Path) -> None:
    ruta = raiz / "nota.txt"
    ruta.write_text("Texto literal, sin transformar.", encoding="utf-8")

    resultado = TxtMarkdownExtractor().extract(str(ruta), FormatoDocumento.TXT, TipoProcesamiento.EXTRACCION_DIRECTA)
    assert resultado.markdown == "Texto literal, sin transformar."
    assert resultado.metodo == "texto_plano"
    assert resultado.advertencias == []


def test_markdown_extractor_lee_literal_sin_reinterpretar(raiz: Path) -> None:
    ruta = raiz / "nota.md"
    ruta.write_text("# Ya es Markdown\n\ncontenido", encoding="utf-8")

    resultado = TxtMarkdownExtractor().extract(str(ruta), FormatoDocumento.MARKDOWN, TipoProcesamiento.EXTRACCION_DIRECTA)
    assert resultado.markdown == "# Ya es Markdown\n\ncontenido"
    assert resultado.metodo == "markdown"


def test_txt_vacio_deja_advertencia(raiz: Path) -> None:
    ruta = raiz / "vacio.txt"
    ruta.write_text("", encoding="utf-8")

    resultado = TxtMarkdownExtractor().extract(str(ruta), FormatoDocumento.TXT, TipoProcesamiento.EXTRACCION_DIRECTA)
    assert resultado.advertencias


def test_json_extractor_produce_bloque_de_codigo(raiz: Path) -> None:
    ruta = raiz / "datos.json"
    ruta.write_text(json.dumps({"area": "rh", "activo": True}), encoding="utf-8")

    resultado = JsonExtractor().extract(str(ruta), FormatoDocumento.JSON, TipoProcesamiento.EXTRACCION_DIRECTA)
    assert "```json" in resultado.markdown
    assert '"area": "rh"' in resultado.markdown
    assert resultado.metodo == "json"
    assert resultado.advertencias == []


def test_json_extractor_vacio_deja_advertencia(raiz: Path) -> None:
    ruta = raiz / "vacio.json"
    ruta.write_text("{}", encoding="utf-8")

    resultado = JsonExtractor().extract(str(ruta), FormatoDocumento.JSON, TipoProcesamiento.EXTRACCION_DIRECTA)
    assert resultado.advertencias


def test_json_extractor_invalido_lanza_excepcion(raiz: Path) -> None:
    ruta = raiz / "invalido.json"
    ruta.write_text("{no es json valido", encoding="utf-8")

    with pytest.raises(ExtractionFailedError):
        JsonExtractor().extract(str(ruta), FormatoDocumento.JSON, TipoProcesamiento.EXTRACCION_DIRECTA)


def test_docx_extractor_normaliza_encabezados_y_tablas(raiz: Path) -> None:
    ruta = raiz / "documento.docx"
    d = docx.Document()
    d.add_heading("Titulo principal", level=1)
    d.add_paragraph("Un parrafo normal.")
    tabla = d.add_table(rows=2, cols=2)
    tabla.rows[0].cells[0].text = "A"
    tabla.rows[0].cells[1].text = "B"
    tabla.rows[1].cells[0].text = "1"
    tabla.rows[1].cells[1].text = "2"
    d.save(str(ruta))

    resultado = DocxExtractor().extract(str(ruta), FormatoDocumento.DOCX, TipoProcesamiento.EXTRACCION_DIRECTA)
    assert "# Titulo principal" in resultado.markdown
    assert "Un parrafo normal." in resultado.markdown
    assert "| A | B |" in resultado.markdown
    assert resultado.metodo == "docx"


def test_docx_extractor_vacio_deja_advertencia(raiz: Path) -> None:
    ruta = raiz / "vacio.docx"
    docx.Document().save(str(ruta))

    resultado = DocxExtractor().extract(str(ruta), FormatoDocumento.DOCX, TipoProcesamiento.EXTRACCION_DIRECTA)
    assert resultado.advertencias


def test_docx_extractor_archivo_invalido_lanza_excepcion(raiz: Path) -> None:
    ruta = raiz / "no_es_docx.docx"
    ruta.write_bytes(b"esto no es un docx valido")

    with pytest.raises(ExtractionFailedError):
        DocxExtractor().extract(str(ruta), FormatoDocumento.DOCX, TipoProcesamiento.EXTRACCION_DIRECTA)


def test_pdf_extractor_extrae_texto_real(raiz: Path) -> None:
    ruta = raiz / "documento.pdf"
    c = canvas.Canvas(str(ruta))
    c.drawString(100, 700, "Contenido de prueba extraible.")
    c.save()

    resultado = PdfExtractor().extract(str(ruta), FormatoDocumento.PDF, TipoProcesamiento.EXTRACCION_DIRECTA)
    assert "Contenido de prueba extraible." in resultado.markdown
    assert resultado.metodo == "pdf_texto"
    assert resultado.advertencias == []


def test_pdf_extractor_pagina_unica_mapa_paginas_empieza_en_cero(raiz: Path) -> None:
    ruta = raiz / "documento.pdf"
    c = canvas.Canvas(str(ruta))
    c.drawString(100, 700, "Contenido de prueba extraible.")
    c.save()

    resultado = PdfExtractor().extract(str(ruta), FormatoDocumento.PDF, TipoProcesamiento.EXTRACCION_DIRECTA)
    assert resultado.mapa_paginas == [0]


def test_pdf_extractor_pagina_sin_texto_deja_advertencia(raiz: Path) -> None:
    from pypdf import PdfWriter

    ruta = raiz / "en_blanco.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    with ruta.open("wb") as f:
        writer.write(f)

    resultado = PdfExtractor().extract(str(ruta), FormatoDocumento.PDF, TipoProcesamiento.EXTRACCION_DIRECTA)
    assert resultado.markdown == ""
    assert resultado.advertencias


def test_pdf_extractor_mapa_paginas_de_varias_paginas_coincide_con_los_offsets_reales(raiz: Path) -> None:
    ruta = raiz / "multipagina.pdf"
    c = canvas.Canvas(str(ruta))
    c.drawString(100, 700, "Primera pagina.")
    c.showPage()
    c.drawString(100, 700, "Segunda pagina con mas texto.")
    c.save()

    resultado = PdfExtractor().extract(str(ruta), FormatoDocumento.PDF, TipoProcesamiento.EXTRACCION_DIRECTA)
    assert len(resultado.mapa_paginas) == 2
    assert resultado.mapa_paginas[0] == 0
    # El offset de la segunda pagina debe apuntar exactamente a donde
    # empieza su texto dentro del markdown final -- mismo invariante que
    # explota la GUI para el salto de cursor.
    segundo_offset = resultado.mapa_paginas[1]
    assert resultado.markdown[segundo_offset:].startswith("Segunda pagina")


def test_pdf_extractor_ilegible_lanza_excepcion(raiz: Path) -> None:
    ruta = raiz / "no_es_pdf.pdf"
    ruta.write_bytes(b"esto no es un pdf")

    with pytest.raises(ExtractionFailedError):
        PdfExtractor().extract(str(ruta), FormatoDocumento.PDF, TipoProcesamiento.EXTRACCION_DIRECTA)


def test_xlsx_extractor_produce_tabla_por_hoja(raiz: Path) -> None:
    ruta = raiz / "libro.xlsx"
    libro = openpyxl.Workbook()
    hoja = libro.active
    hoja.title = "Datos"
    hoja.append(["nombre", "area"])
    hoja.append(["Ana", "RH"])
    libro.save(str(ruta))

    resultado = XlsxExtractor().extract(str(ruta), FormatoDocumento.XLSX, TipoProcesamiento.EXTRACCION_DIRECTA)
    assert "## Datos" in resultado.markdown
    assert "| nombre | area |" in resultado.markdown
    assert "| Ana | RH |" in resultado.markdown


def test_xlsx_extractor_hoja_vacia_deja_advertencia(raiz: Path) -> None:
    ruta = raiz / "vacio.xlsx"
    openpyxl.Workbook().save(str(ruta))

    resultado = XlsxExtractor().extract(str(ruta), FormatoDocumento.XLSX, TipoProcesamiento.EXTRACCION_DIRECTA)
    assert resultado.advertencias


def test_csv_extractor_produce_tabla_markdown(raiz: Path) -> None:
    ruta = raiz / "datos.csv"
    ruta.write_text("nombre,area\nAna,RH\nLuis,Legal\n", encoding="utf-8")

    resultado = CsvExtractor().extract(str(ruta), FormatoDocumento.CSV, TipoProcesamiento.EXTRACCION_DIRECTA)
    assert "| nombre | area |" in resultado.markdown
    assert "| Ana | RH |" in resultado.markdown


def test_csv_extractor_vacio_deja_advertencia(raiz: Path) -> None:
    ruta = raiz / "vacio.csv"
    ruta.write_text("", encoding="utf-8")

    resultado = CsvExtractor().extract(str(ruta), FormatoDocumento.CSV, TipoProcesamiento.EXTRACCION_DIRECTA)
    assert resultado.markdown == ""
    assert resultado.advertencias


def test_composite_extractor_despacha_por_formato(raiz: Path) -> None:
    ruta = raiz / "nota.txt"
    ruta.write_text("hola", encoding="utf-8")

    resultado = CompositeDocumentExtractor().extract(str(ruta), FormatoDocumento.TXT, TipoProcesamiento.EXTRACCION_DIRECTA)
    assert resultado.metodo == "texto_plano"


def test_composite_extractor_formato_no_soportado_lanza_excepcion(raiz: Path) -> None:
    ruta = raiz / "imagen.png"
    ruta.write_bytes(b"contenido binario simulado")

    with pytest.raises(UnsupportedFormatError):
        CompositeDocumentExtractor().extract(str(ruta), FormatoDocumento.PNG, TipoProcesamiento.EXTRACCION_DIRECTA)


def test_txt_con_bom_de_windows_no_deja_el_bom_en_el_markdown(raiz: Path) -> None:
    """Auditoria 2026-09-30: un .txt/.md guardado con BOM (Bloc de notas de
    Windows) dejaba U+FEFF al inicio del Markdown y el "# Titulo" inicial
    ya no era un encabezado."""
    ruta = raiz / "con_bom.md"
    ruta.write_bytes(b"\xef\xbb\xbf# Titulo\r\n\r\nTexto.\r\n")

    resultado = TxtMarkdownExtractor().extract(str(ruta), FormatoDocumento.MARKDOWN, TipoProcesamiento.EXTRACCION_DIRECTA)

    assert resultado.markdown.startswith("# Titulo")
    assert "﻿" not in resultado.markdown


def test_json_con_bom_de_windows_se_extrae_sin_error(raiz: Path) -> None:
    """Auditoria 2026-09-30: antes fallaba con "JSON invalido: Unexpected UTF-8 BOM"."""
    ruta = raiz / "con_bom.json"
    ruta.write_bytes(b"\xef\xbb\xbf" + json.dumps({"carrera": "Derecho"}).encode("utf-8"))

    resultado = JsonExtractor().extract(str(ruta), FormatoDocumento.JSON, TipoProcesamiento.EXTRACCION_DIRECTA)

    assert '"carrera": "Derecho"' in resultado.markdown
