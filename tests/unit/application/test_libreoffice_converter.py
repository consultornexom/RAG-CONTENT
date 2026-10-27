"""Prueba real de LibreOfficeConverter contra el binario 'soffice'
instalado en este entorno -- no un doble de prueba: el proposito de
este adaptador es justamente envolver un proceso externo real. Los
archivos DOC/XLS de prueba se generan con el propio LibreOffice
(round-trip DOCX->DOC y XLSX->XLS) porque no hay una forma simple de
producir un binario OLE2 valido sin el.
"""
from __future__ import annotations

import shutil
import subprocess
import tempfile
from pathlib import Path

import docx
import openpyxl
import pytest

from domain.documents.entities import FormatoDocumento
from domain.documents.exceptions import LegacyConversionFailedError
from infrastructure.extraction.libreoffice_converter import LibreOfficeConverter

RAIZ_PRUEBA = Path("/tmp/rag_content_studio_libreoffice_converter_test")


@pytest.fixture()
def raiz():
    if RAIZ_PRUEBA.exists():
        shutil.rmtree(RAIZ_PRUEBA)
    RAIZ_PRUEBA.mkdir(parents=True)
    yield RAIZ_PRUEBA
    shutil.rmtree(RAIZ_PRUEBA, ignore_errors=True)


def _convertir_con_soffice(ruta_entrada: Path, filtro_destino: str, directorio_salida: Path) -> None:
    perfil = tempfile.mkdtemp(prefix="rag_content_studio_test_lo_profile_")
    try:
        subprocess.run(
            [
                "soffice",
                "--headless",
                "--norestore",
                f"-env:UserInstallation=file://{perfil}",
                "--convert-to",
                filtro_destino,
                "--outdir",
                str(directorio_salida),
                str(ruta_entrada),
            ],
            capture_output=True,
            text=True,
            timeout=120,
            check=True,
        )
    finally:
        shutil.rmtree(perfil, ignore_errors=True)


def _doc_de_prueba(raiz: Path) -> Path:
    d = docx.Document()
    d.add_heading("Contrato legado", level=1)
    d.add_paragraph("Texto de prueba en formato DOC clasico.")
    docx_path = raiz / "contrato.docx"
    d.save(str(docx_path))
    _convertir_con_soffice(docx_path, "doc", raiz)
    return raiz / "contrato.doc"


def _xls_de_prueba(raiz: Path) -> Path:
    libro = openpyxl.Workbook()
    hoja = libro.active
    hoja.title = "Datos"
    hoja.append(["nombre", "area"])
    hoja.append(["Ana", "RH"])
    xlsx_path = raiz / "datos.xlsx"
    libro.save(str(xlsx_path))
    _convertir_con_soffice(xlsx_path, "xls", raiz)
    return raiz / "datos.xls"


def test_convertir_doc_a_docx_produce_un_docx_valido(raiz: Path) -> None:
    ruta_doc = _doc_de_prueba(raiz)

    ruta_convertida = LibreOfficeConverter().convertir(str(ruta_doc), FormatoDocumento.DOCX)
    try:
        assert ruta_convertida.endswith(".docx")
        documento = docx.Document(ruta_convertida)
        texto_completo = "\n".join(p.text for p in documento.paragraphs)
        assert "Contrato legado" in texto_completo
        assert "Texto de prueba en formato DOC clasico." in texto_completo
    finally:
        shutil.rmtree(Path(ruta_convertida).parent, ignore_errors=True)


def test_convertir_xls_a_xlsx_produce_un_xlsx_valido(raiz: Path) -> None:
    ruta_xls = _xls_de_prueba(raiz)

    ruta_convertida = LibreOfficeConverter().convertir(str(ruta_xls), FormatoDocumento.XLSX)
    try:
        assert ruta_convertida.endswith(".xlsx")
        libro = openpyxl.load_workbook(ruta_convertida, read_only=True)
        filas = list(libro.active.iter_rows(values_only=True))
        assert filas[0] == ("nombre", "area")
        assert filas[1] == ("Ana", "RH")
    finally:
        shutil.rmtree(Path(ruta_convertida).parent, ignore_errors=True)


def test_convertir_archivo_inexistente_lanza_excepcion_clara(raiz: Path) -> None:
    with pytest.raises(LegacyConversionFailedError):
        LibreOfficeConverter().convertir(str(raiz / "no_existe.doc"), FormatoDocumento.DOCX)


def test_convertir_a_formato_destino_no_soportado_lanza_excepcion_clara(raiz: Path) -> None:
    ruta_doc = _doc_de_prueba(raiz)
    with pytest.raises(LegacyConversionFailedError, match="destino"):
        LibreOfficeConverter().convertir(str(ruta_doc), FormatoDocumento.PDF)
