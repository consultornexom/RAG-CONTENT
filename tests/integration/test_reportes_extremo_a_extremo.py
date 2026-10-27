"""Prueba de humo de la Fase 8 (Reportes y documentacion automatica).
Mismo patron que test_cerrar_lote_extremo_a_extremo.py: bootstrap()/
abrir_proyecto() reales, sin fakes, para probar el cableado completo de
app.bootstrap (ReportsService sobre los repositorios reales + 10_reportes/
del proyecto)."""
from __future__ import annotations

import shutil
from pathlib import Path

import docx
import pytest

from app.bootstrap import abrir_proyecto, bootstrap
from domain.incidents.entities import EtapaIncidencia, SeveridadIncidencia

RAIZ_PRUEBA = Path("/tmp/rag_content_studio_smoke_test_reportes")


@pytest.fixture()
def raiz_temporal():
    if RAIZ_PRUEBA.exists():
        shutil.rmtree(RAIZ_PRUEBA)
    RAIZ_PRUEBA.mkdir(parents=True)
    perfiles_dir = RAIZ_PRUEBA / "perfiles" / "perfil_generico"
    perfiles_dir.mkdir(parents=True)
    shutil.copy(
        Path(__file__).resolve().parents[2] / "perfiles" / "perfil_generico" / "0.1.yaml",
        perfiles_dir / "0.1.yaml",
    )
    origenes_dir = RAIZ_PRUEBA / "_origenes_de_prueba"
    origenes_dir.mkdir()
    yield RAIZ_PRUEBA, origenes_dir
    shutil.rmtree(RAIZ_PRUEBA, ignore_errors=True)


def test_los_siete_documentos_de_la_seccion_7_1_se_generan_de_verdad(raiz_temporal) -> None:
    raiz, origenes_dir = raiz_temporal
    ctx = bootstrap(raiz=str(raiz))
    proyecto = ctx.projects.crear_proyecto(nombre="Fase 8", perfil_id="perfil_generico")
    pctx = abrir_proyecto(ctx, proyecto)

    lote = pctx.batches.crear_lote(proyecto.id, "Lote de enero")

    # Documento 1: flujo completo real hasta LIBERADO.
    archivo_1 = origenes_dir / "politica.docx"
    d = docx.Document()
    d.add_heading("Politica de vacaciones", level=1)
    d.add_paragraph("Contenido real de la politica.")
    d.save(str(archivo_1))

    doc_1 = pctx.documents.registrar_documento(proyecto_id=proyecto.id, ruta_origen=str(archivo_1))
    doc_1 = pctx.documents.clasificar_documento(doc_1.id)
    doc_1 = pctx.documents.extraer_documento(doc_1.id)
    doc_1 = pctx.documents.revisar_documento(doc_1.id, "ana@nexom.mx", "se ve bien")
    doc_1 = pctx.documents.normalizar_documento(doc_1.id)
    doc_1 = pctx.documents.aprobar_para_markdown(doc_1.id, "ana@nexom.mx")
    doc_1 = pctx.documents.generar_markdown(doc_1.id)
    doc_1 = pctx.documents.validar_markdown(doc_1.id)
    pctx.validity.registrar_fuente(doc_1.id, "portal interno", "RH", "Confirmado el 2026-07-28.")
    pctx.batches.agregar_documento(lote.id, doc_1.id)

    # Documento 2: se deja apenas registrado -- CerrarLote lo bloquea.
    archivo_2 = origenes_dir / "memo.txt"
    archivo_2.write_text("Memo corto sin procesar todavia.", encoding="utf-8")
    doc_2 = pctx.documents.registrar_documento(proyecto_id=proyecto.id, ruta_origen=str(archivo_2))
    pctx.batches.agregar_documento(lote.id, doc_2.id)

    incidencia = pctx.incidents.registrar_incidencia(
        proyecto_id=proyecto.id,
        documento_id=doc_1.id,
        etapa=EtapaIncidencia.REVISION,
        tipo="cifra_dudosa",
        severidad=SeveridadIncidencia.MEDIA,
        evidencia="parrafo 2",
        impacto="posible dato desactualizado",
    )
    assert incidencia is not None

    lote_cerrado = pctx.batches.cerrar_lote(lote.id)
    assert lote_cerrado.estado.value == "cerrado"

    rutas = pctx.reports.generar_todos_los_informes()
    assert set(rutas) == {"inventario", "procesamiento", "revision", "incidencias", "calidad", "trazabilidad"}
    for ruta in rutas.values():
        assert Path(ruta).parent.name == "10_reportes"
        assert Path(ruta).exists()

    contenido_inventario = Path(rutas["inventario"]).read_text(encoding="utf-8")
    assert "politica.docx" in contenido_inventario
    assert "memo.txt" in contenido_inventario

    contenido_calidad = Path(rutas["calidad"]).read_text(encoding="utf-8")
    assert "liberado" in contenido_calidad  # doc_1
    assert "bloqueado" in contenido_calidad  # doc_2

    contenido_incidencias = Path(rutas["incidencias"]).read_text(encoding="utf-8")
    assert "cifra_dudosa" in contenido_incidencias

    contenido_trazabilidad = Path(rutas["trazabilidad"]).read_text(encoding="utf-8")
    assert "ana@nexom.mx" in contenido_trazabilidad

    ruta_acta = pctx.reports.generar_acta_cierre_lote(lote.id)
    assert Path(ruta_acta).parent.name == "10_reportes"
    contenido_acta = Path(ruta_acta).read_text(encoding="utf-8")
    assert doc_1.id in contenido_acta
    assert doc_2.id in contenido_acta
    assert "ana@nexom.mx" in contenido_acta
