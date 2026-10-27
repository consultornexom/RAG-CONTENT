"""Prueba de humo del Modulo 12 (Generador de Markdown) y de la tercera
regla del Modulo 13 (no liberar sin fuente y aprobacion). Usa una raiz
temporal, nunca la ruta real de datos -- mismo patron que
test_normalizar_documento_extremo_a_extremo.py.

Modulo 8 (revision) todavia no tiene un paso de GUI para este tramo del
flujo -- se avanza NORMALIZADO con los mismos primitivos ya usados por
la prueba de humo de Modulo 9 (avanzar_a() directo sobre el repositorio
real, no un fake) hasta el punto donde este modulo empieza a actuar de
verdad: aprobar_para_markdown() en adelante ya es el caso de uso real
de DocumentsService/GovernanceService, cableado por app.bootstrap tal
cual lo usaria la interfaz."""
from __future__ import annotations

import shutil
from pathlib import Path

import docx
import pytest

from app.bootstrap import abrir_proyecto, bootstrap
from domain.documents.entities import EstadoDocumento
from domain.governance.exceptions import AprobacionFaltanteError, FuenteFaltanteError
from infrastructure.database.document_repository import SqlAlchemyDocumentRepository

RAIZ_PRUEBA = Path("/tmp/rag_content_studio_smoke_test_markdown_final")


@pytest.fixture()
def raiz_temporal():
    if RAIZ_PRUEBA.exists():
        shutil.rmtree(RAIZ_PRUEBA)
    RAIZ_PRUEBA.mkdir(parents=True)
    perfiles_dir = RAIZ_PRUEBA / "perfiles" / "perfil_generico"
    perfiles_dir.mkdir(parents=True)
    shutil.copy(
        Path(__file__).resolve().parents[2] / "perfiles" / "perfil_generico" / "0.1.yaml",
        perfiles_dir / "0.1.yaml",
    )
    origenes_dir = RAIZ_PRUEBA / "_origenes_de_prueba"
    origenes_dir.mkdir()
    yield RAIZ_PRUEBA, origenes_dir
    shutil.rmtree(RAIZ_PRUEBA, ignore_errors=True)


def _avanzar_a_revisado(pctx, documento_id: str):
    """Sustituto de Modulo 8 (todavia sin GUI para este tramo): usa el
    mismo repositorio real (SqlAlchemy, sobre la sesion ya abierta de
    proyecto.db) que ProjectContext, no un fake -- mismo patron que
    test_normalizar_documento_extremo_a_extremo.py."""
    repo = SqlAlchemyDocumentRepository(pctx.session)
    documento = repo.get(documento_id)
    assert documento is not None
    revisado = documento.avanzar_a(EstadoDocumento.EN_REVISION).avanzar_a(EstadoDocumento.REVISADO)
    repo.save(revisado)
    return revisado


def test_flujo_completo_hasta_liberado_y_manifiesto(raiz_temporal) -> None:
    raiz, origenes_dir = raiz_temporal
    ctx = bootstrap(raiz=str(raiz))
    proyecto = ctx.projects.crear_proyecto(nombre="Modulo 12", perfil_id="perfil_generico")
    pctx = abrir_proyecto(ctx, proyecto)

    archivo = origenes_dir / "politica.docx"
    d = docx.Document()
    d.add_heading("Politica de vacaciones", level=1)
    d.add_paragraph("Contenido real de la politica, ya revisado y normalizado.")
    d.save(str(archivo))

    documento = pctx.documents.registrar_documento(
        proyecto_id=proyecto.id, ruta_origen=str(archivo), metadatos={"titulo_oficial": "Politica de vacaciones"}
    )
    clasificado = pctx.documents.clasificar_documento(documento.id)
    extraido = pctx.documents.extraer_documento(clasificado.id)
    _avanzar_a_revisado(pctx, extraido.id)
    normalizado = pctx.documents.normalizar_documento(extraido.id)
    assert normalizado.estado == EstadoDocumento.NORMALIZADO

    # -- Aprobar para markdown: primer punto de control del Modulo 13.
    aprobado = pctx.documents.aprobar_para_markdown(normalizado.id, "ana@nexom.mx", "Listo para generar markdown.")
    assert aprobado.estado == EstadoDocumento.APROBADO_PARA_MARKDOWN
    aprobaciones = pctx.documents.listar_aprobaciones(normalizado.id)
    assert len(aprobaciones) == 1
    assert aprobaciones[0].aprobador == "ana@nexom.mx"

    # -- Generar el Markdown final: debe quedar en 07_markdown/, no en
    # 06_normalizados/.
    generado = pctx.documents.generar_markdown(normalizado.id)
    assert generado.estado == EstadoDocumento.MARKDOWN_GENERADO
    assert Path(generado.ruta_markdown).parent == Path(proyecto.ruta_carpeta) / "07_markdown"
    assert generado.hash_markdown is not None
    contenido_final = Path(generado.ruta_markdown).read_text(encoding="utf-8")
    assert "Politica de vacaciones" in contenido_final

    # -- Validar: el documento tiene un encabezado real, sin advertencias.
    validado = pctx.documents.validar_markdown(normalizado.id)
    assert validado.estado == EstadoDocumento.VALIDADO
    assert validado.advertencias_markdown == []

    # -- Liberar sin fuente registrada: debe bloquear (tercera regla del
    # Modulo 13, "no liberar contenido sin fuente, hash y aprobacion").
    with pytest.raises(FuenteFaltanteError):
        pctx.documents.liberar_documento(normalizado.id)
    assert pctx.documents.obtener_documento(normalizado.id).estado == EstadoDocumento.VALIDADO

    # -- Registrar la fuente (Modulo 11) y liberar de verdad.
    pctx.validity.registrar_fuente(
        normalizado.id, "portal de normatividad interna", "Direccion Juridica", "Confirmado el 2026-07-28."
    )
    liberado = pctx.documents.liberar_documento(normalizado.id)
    assert liberado.estado == EstadoDocumento.LIBERADO

    # -- Generar el manifiesto: debe incluir el documento recien liberado,
    # con el hash del Markdown final (no el hash_sha256 del original).
    ruta_manifiesto = pctx.documents.generar_manifiesto()
    assert Path(ruta_manifiesto).parent == Path(proyecto.ruta_carpeta) / "12_exportaciones"
    contenido_manifiesto = Path(ruta_manifiesto).read_text(encoding="utf-8")
    assert "Politica de vacaciones" in contenido_manifiesto
    assert liberado.hash_markdown in contenido_manifiesto
    assert liberado.hash_sha256 not in contenido_manifiesto


def test_liberar_sin_aprobacion_falla_incluso_con_fuente_registrada(raiz_temporal) -> None:
    """No alcanzable por el flujo normal de la GUI (aprobar_para_markdown
    siempre crea una Aprobacion antes de generar el markdown), pero
    confirma que la regla de liberacion revisa AMBOS requisitos por
    separado, no solo el de fuente."""
    raiz, origenes_dir = raiz_temporal
    ctx = bootstrap(raiz=str(raiz))
    proyecto = ctx.projects.crear_proyecto(nombre="Modulo 12 sin aprobacion", perfil_id="perfil_generico")
    pctx = abrir_proyecto(ctx, proyecto)

    archivo = origenes_dir / "memo.txt"
    archivo.write_text("Memo corto sin estructura Markdown.", encoding="utf-8")

    documento = pctx.documents.registrar_documento(proyecto_id=proyecto.id, ruta_origen=str(archivo))
    clasificado = pctx.documents.clasificar_documento(documento.id)
    extraido = pctx.documents.extraer_documento(clasificado.id)
    _avanzar_a_revisado(pctx, extraido.id)
    normalizado = pctx.documents.normalizar_documento(extraido.id)

    # Bypass deliberado del paso real de aprobacion (avanza el estado
    # directo sobre el repositorio) para poder llegar a VALIDADO sin
    # ninguna Aprobacion registrada.
    repo = SqlAlchemyDocumentRepository(pctx.session)
    forzado = normalizado.avanzar_a(EstadoDocumento.PENDIENTE_DE_APROBACION).avanzar_a(
        EstadoDocumento.APROBADO_PARA_MARKDOWN
    )
    repo.save(forzado)

    generado = pctx.documents.generar_markdown(normalizado.id)
    pctx.documents.validar_markdown(generado.id)
    pctx.validity.registrar_fuente(normalizado.id, "correo interno", "RH", "Recibido por correo el 2026-07-28.")

    with pytest.raises(AprobacionFaltanteError):
        pctx.documents.liberar_documento(normalizado.id)
