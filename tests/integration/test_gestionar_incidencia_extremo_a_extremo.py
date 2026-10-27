"""Prueba de humo del Modulo 10 (gestion de incidencias). Usa una raiz
temporal, nunca la ruta real de datos -- mismo patron que
test_extraer_documento_extremo_a_extremo.py y
test_normalizar_documento_extremo_a_extremo.py.
"""
from __future__ import annotations

import shutil
from pathlib import Path

import docx
import pytest

from app.bootstrap import abrir_proyecto, bootstrap
from domain.incidents.entities import EstadoIncidencia, EtapaIncidencia, SeveridadIncidencia

RAIZ_PRUEBA = Path("/tmp/rag_content_studio_smoke_test_incidencias")


@pytest.fixture()
def raiz_temporal():
    if RAIZ_PRUEBA.exists():
        shutil.rmtree(RAIZ_PRUEBA)
    RAIZ_PRUEBA.mkdir(parents=True)
    perfiles_dir = RAIZ_PRUEBA / "perfiles" / "perfil_generico"
    perfiles_dir.mkdir(parents=True)
    shutil.copy(
        Path(__file__).resolve().parents[2] / "perfiles" / "perfil_generico" / "0.1.yaml",
        perfiles_dir / "0.1.yaml",
    )
    origenes_dir = RAIZ_PRUEBA / "_origenes_de_prueba"
    origenes_dir.mkdir()
    yield RAIZ_PRUEBA, origenes_dir
    shutil.rmtree(RAIZ_PRUEBA, ignore_errors=True)


def test_registrar_y_resolver_una_incidencia_sobre_un_documento_real(raiz_temporal) -> None:
    raiz, origenes_dir = raiz_temporal
    ctx = bootstrap(raiz=str(raiz))
    proyecto = ctx.projects.crear_proyecto(nombre="Incidencias", perfil_id="perfil_generico")
    pctx = abrir_proyecto(ctx, proyecto)

    archivo = origenes_dir / "contrato.docx"
    d = docx.Document()
    d.add_heading("Contrato de prueba", level=1)
    d.add_paragraph("El monto total es de 1,000,000 MXN.")
    d.save(str(archivo))

    documento = pctx.documents.registrar_documento(proyecto_id=proyecto.id, ruta_origen=str(archivo))
    pctx.documents.clasificar_documento(documento.id)
    extraido = pctx.documents.extraer_documento(documento.id)

    incidencia = pctx.incidents.registrar_incidencia(
        proyecto_id=proyecto.id,
        documento_id=extraido.id,
        etapa=EtapaIncidencia.EXTRACCION,
        tipo="cifra_dudosa",
        severidad=SeveridadIncidencia.CRITICA,
        evidencia="El monto '1,000,000 MXN' no coincide con el original consultado por el revisor.",
        impacto="Podria invalidar el monto total reportado en el Markdown final.",
    )
    assert incidencia.estado == EstadoIncidencia.ABIERTA
    assert pctx.incidents.existen_criticas_abiertas(extraido.id)

    pctx.incidents.asignar_incidencia(incidencia.id, "ana@nexom.mx")
    pctx.incidents.responder_incidencia(incidencia.id, "Se coteja contra el original: la cifra si es correcta.")
    resuelta = pctx.incidents.decidir_incidencia(
        incidencia.id, "Confirmado correcto contra el original, sin cambios necesarios.", resuelta=True
    )
    assert resuelta.estado == EstadoIncidencia.RESUELTA
    # RESUELTA todavia no es CERRADA -- sigue contando como abierta
    # para la regla de bloqueo de liberacion (Documento Maestro, seccion
    # 8.2), responsabilidad real del Modulo 13.
    assert pctx.incidents.existen_criticas_abiertas(extraido.id)

    cerrada = pctx.incidents.cerrar_incidencia(incidencia.id, "Verificado y cerrado tras cotejar el original.")
    assert cerrada.estado == EstadoIncidencia.CERRADA
    assert not pctx.incidents.existen_criticas_abiertas(extraido.id)

    # Recuperada de proyecto.db, no el objeto en memoria.
    recuperada = pctx.incidents.obtener_incidencia(incidencia.id)
    assert recuperada is not None
    assert recuperada.estado == EstadoIncidencia.CERRADA
    assert recuperada.responsable == "ana@nexom.mx"
    assert recuperada.cierre == "Verificado y cerrado tras cotejar el original."


def test_registrar_incidencia_sobre_documento_inexistente_falla(raiz_temporal) -> None:
    raiz, _origenes_dir = raiz_temporal
    ctx = bootstrap(raiz=str(raiz))
    proyecto = ctx.projects.crear_proyecto(nombre="Incidencias sin documento", perfil_id="perfil_generico")
    pctx = abrir_proyecto(ctx, proyecto)

    from domain.documents.exceptions import DocumentNotFoundError

    with pytest.raises(DocumentNotFoundError):
        pctx.incidents.registrar_incidencia(
            proyecto_id=proyecto.id,
            documento_id="no-existe",
            etapa=EtapaIncidencia.REVISION,
            tipo="cifra_dudosa",
            severidad=SeveridadIncidencia.ALTA,
            evidencia="evidencia",
            impacto="impacto",
        )
