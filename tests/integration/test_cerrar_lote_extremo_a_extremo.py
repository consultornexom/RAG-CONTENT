"""Prueba de humo de la Fase 7 (Lotes). Usa una raiz temporal, nunca la
ruta real de datos -- mismo patron que el resto de
tests/smoke/test_*_extremo_a_extremo.py: bootstrap()/abrir_proyecto()
reales, sin fakes, para probar el cableado completo de app.bootstrap
(BatchesService sobre LoteRepository/DocumentRepository/GovernanceService
reales)."""
from __future__ import annotations

import shutil
from pathlib import Path

import docx
import pytest

from app.bootstrap import abrir_proyecto, bootstrap
from domain.documents.entities import EstadoDocumento
from infrastructure.database.document_repository import SqlAlchemyDocumentRepository

RAIZ_PRUEBA = Path("/tmp/rag_content_studio_smoke_test_lotes")


@pytest.fixture()
def raiz_temporal():
    if RAIZ_PRUEBA.exists():
        shutil.rmtree(RAIZ_PRUEBA)
    RAIZ_PRUEBA.mkdir(parents=True)
    perfiles_dir = RAIZ_PRUEBA / "perfiles" / "perfil_generico"
    perfiles_dir.mkdir(parents=True)
    shutil.copy(
        Path(__file__).resolve().parents[2] / "perfiles" / "perfil_generico" / "0.1.yaml",
        perfiles_dir / "0.1.yaml",
    )
    origenes_dir = RAIZ_PRUEBA / "_origenes_de_prueba"
    origenes_dir.mkdir()
    yield RAIZ_PRUEBA, origenes_dir
    shutil.rmtree(RAIZ_PRUEBA, ignore_errors=True)


def _avanzar_a_revisado(pctx, documento_id: str):
    """Sustituto de Modulo 8 (todavia sin GUI para este tramo): mismo
    patron que los demas tests de humo posteriores al Modulo 8."""
    repo = SqlAlchemyDocumentRepository(pctx.session)
    documento = repo.get(documento_id)
    assert documento is not None
    revisado = documento.avanzar_a(EstadoDocumento.EN_REVISION).avanzar_a(EstadoDocumento.REVISADO)
    repo.save(revisado)
    return revisado


def test_cerrar_lote_libera_el_aprobado_y_bloquea_el_pendiente(raiz_temporal) -> None:
    raiz, origenes_dir = raiz_temporal
    ctx = bootstrap(raiz=str(raiz))
    proyecto = ctx.projects.crear_proyecto(nombre="Fase 7", perfil_id="perfil_generico")
    pctx = abrir_proyecto(ctx, proyecto)

    lote = pctx.batches.crear_lote(proyecto.id, "Lote de enero")

    # -- Documento 1: se lleva hasta VALIDADO, con Fuente y Aprobacion
    # reales -- CerrarLote debe liberarlo de verdad.
    archivo_1 = origenes_dir / "politica.docx"
    d = docx.Document()
    d.add_heading("Politica de vacaciones", level=1)
    d.add_paragraph("Contenido real de la politica.")
    d.save(str(archivo_1))

    doc_1 = pctx.documents.registrar_documento(proyecto_id=proyecto.id, ruta_origen=str(archivo_1))
    doc_1 = pctx.documents.clasificar_documento(doc_1.id)
    doc_1 = pctx.documents.extraer_documento(doc_1.id)
    _avanzar_a_revisado(pctx, doc_1.id)
    doc_1 = pctx.documents.normalizar_documento(doc_1.id)
    doc_1 = pctx.documents.aprobar_para_markdown(doc_1.id, "ana@nexom.mx")
    doc_1 = pctx.documents.generar_markdown(doc_1.id)
    doc_1 = pctx.documents.validar_markdown(doc_1.id)
    pctx.validity.registrar_fuente(doc_1.id, "portal interno", "RH", "Confirmado el 2026-07-28.")
    pctx.batches.agregar_documento(lote.id, doc_1.id)

    # -- Documento 2: se deja apenas registrado (RECIBIDO) -- CerrarLote
    # debe bloquearlo, nunca liberarlo.
    archivo_2 = origenes_dir / "memo.txt"
    archivo_2.write_text("Memo corto sin procesar todavia.", encoding="utf-8")
    doc_2 = pctx.documents.registrar_documento(proyecto_id=proyecto.id, ruta_origen=str(archivo_2))
    pctx.batches.agregar_documento(lote.id, doc_2.id)

    lote_cerrado = pctx.batches.cerrar_lote(lote.id)

    assert lote_cerrado.estado.value == "cerrado"
    doc_1_final = pctx.documents.obtener_documento(doc_1.id)
    doc_2_final = pctx.documents.obtener_documento(doc_2.id)
    assert doc_1_final.estado == EstadoDocumento.LIBERADO
    assert doc_2_final.estado == EstadoDocumento.BLOQUEADO

    # El manifiesto filtrado por este lote solo debe incluir doc_1.
    ruta_manifiesto = pctx.documents.generar_manifiesto(lote_id=lote.id)
    contenido = Path(ruta_manifiesto).read_text(encoding="utf-8")
    assert "politica.docx" in contenido
    assert "memo.txt" not in contenido  # doc_2 nunca llego a LIBERADO -- no aparece en el manifiesto
