"""Prueba de humo del Modulo 5 (extraccion documental). Usa una raiz
temporal, nunca la ruta real de datos.
"""
from __future__ import annotations

import json
import shutil
from pathlib import Path

import docx
import pytest
from reportlab.pdfgen import canvas

from app.bootstrap import abrir_proyecto, bootstrap

RAIZ_PRUEBA = Path("/tmp/rag_content_studio_smoke_test_extraccion")


@pytest.fixture()
def raiz_temporal():
    if RAIZ_PRUEBA.exists():
        shutil.rmtree(RAIZ_PRUEBA)
    RAIZ_PRUEBA.mkdir(parents=True)
    perfiles_dir = RAIZ_PRUEBA / "perfiles" / "perfil_generico"
    perfiles_dir.mkdir(parents=True)
    shutil.copy(
        Path(__file__).resolve().parents[2] / "perfiles" / "perfil_generico" / "0.1.yaml",
        perfiles_dir / "0.1.yaml",
    )
    origenes_dir = RAIZ_PRUEBA / "_origenes_de_prueba"
    origenes_dir.mkdir()
    yield RAIZ_PRUEBA, origenes_dir
    shutil.rmtree(RAIZ_PRUEBA, ignore_errors=True)


def _registrar_clasificar_y_extraer(pctx, proyecto_id: str, ruta_origen: str):
    documento = pctx.documents.registrar_documento(proyecto_id=proyecto_id, ruta_origen=ruta_origen)
    pctx.documents.clasificar_documento(documento.id)
    return pctx.documents.extraer_documento(documento.id)


def test_extraer_docx_produce_markdown_con_encabezado_y_tabla(raiz_temporal) -> None:
    raiz, origenes_dir = raiz_temporal
    ctx = bootstrap(raiz=str(raiz))
    proyecto = ctx.projects.crear_proyecto(nombre="Extraccion DOCX", perfil_id="perfil_generico")
    pctx = abrir_proyecto(ctx, proyecto)

    archivo = origenes_dir / "contrato.docx"
    d = docx.Document()
    d.add_heading("Contrato de prueba", level=1)
    d.add_paragraph("Texto introductorio.")
    tabla = d.add_table(rows=2, cols=2)
    tabla.rows[0].cells[0].text = "Clausula"
    tabla.rows[0].cells[1].text = "Descripcion"
    tabla.rows[1].cells[0].text = "1"
    tabla.rows[1].cells[1].text = "Vigencia"
    d.save(str(archivo))

    extraido = _registrar_clasificar_y_extraer(pctx, proyecto.id, str(archivo))
    assert extraido.estado.value == "extraido"

    markdown = Path(extraido.ruta_extraccion).read_text(encoding="utf-8")
    assert "# Contrato de prueba" in markdown
    assert "| Clausula | Descripcion |" in markdown

    # Recuperado de proyecto.db, no el objeto en memoria.
    recuperado = pctx.documents.obtener_documento(extraido.id)
    assert recuperado is not None
    assert recuperado.ruta_extraccion == extraido.ruta_extraccion
    assert Path(recuperado.ruta_extraccion).parent == Path(proyecto.ruta_carpeta) / "03_extracciones"


def test_extraer_pdf_con_texto_real(raiz_temporal) -> None:
    raiz, origenes_dir = raiz_temporal
    ctx = bootstrap(raiz=str(raiz))
    proyecto = ctx.projects.crear_proyecto(nombre="Extraccion PDF", perfil_id="perfil_generico")
    pctx = abrir_proyecto(ctx, proyecto)

    archivo = origenes_dir / "politica.pdf"
    c = canvas.Canvas(str(archivo))
    c.drawString(100, 700, "Politica de vacaciones -- texto real extraible.")
    c.save()

    extraido = _registrar_clasificar_y_extraer(pctx, proyecto.id, str(archivo))
    markdown = Path(extraido.ruta_extraccion).read_text(encoding="utf-8")
    assert "Politica de vacaciones" in markdown


def test_extraer_csv_produce_tabla_markdown(raiz_temporal) -> None:
    raiz, origenes_dir = raiz_temporal
    ctx = bootstrap(raiz=str(raiz))
    proyecto = ctx.projects.crear_proyecto(nombre="Extraccion CSV", perfil_id="perfil_generico")
    pctx = abrir_proyecto(ctx, proyecto)

    archivo = origenes_dir / "datos.csv"
    archivo.write_text("nombre,area\nAna,RH\nLuis,Legal\n", encoding="utf-8")

    extraido = _registrar_clasificar_y_extraer(pctx, proyecto.id, str(archivo))
    markdown = Path(extraido.ruta_extraccion).read_text(encoding="utf-8")
    assert "| nombre | area |" in markdown
    assert "| Ana | RH |" in markdown


def test_extraer_json_produce_bloque_de_codigo(raiz_temporal) -> None:
    raiz, origenes_dir = raiz_temporal
    ctx = bootstrap(raiz=str(raiz))
    proyecto = ctx.projects.crear_proyecto(nombre="Extraccion JSON", perfil_id="perfil_generico")
    pctx = abrir_proyecto(ctx, proyecto)

    archivo = origenes_dir / "datos.json"
    archivo.write_text(json.dumps({"clave": "valor"}), encoding="utf-8")

    extraido = _registrar_clasificar_y_extraer(pctx, proyecto.id, str(archivo))
    markdown = Path(extraido.ruta_extraccion).read_text(encoding="utf-8")
    assert "```json" in markdown
    assert '"clave": "valor"' in markdown


def test_extraer_imagen_por_ocr_produce_markdown_con_texto_reconocido(raiz_temporal) -> None:
    """Modulo 6: la ruta OCR ya no se rechaza con
    UnsupportedExtractionRouteError -- una imagen PNG/JPG se clasifica a
    OCR (Modulo 4) y ahora ImageOcrExtractor + TesseractOcrEngine la
    procesan de verdad, cableados por app.bootstrap.abrir_proyecto().

    El idioma por defecto cableado es 'spa+eng' (ADR-001, seccion
    14.1); este entorno de pruebas solo tiene instalado el paquete
    'eng' de Tesseract (no 'spa'), pero eso no produce
    OcrEngineUnavailableError: Tesseract ignora en silencio un idioma
    solicitado sin paquete instalado siempre que al menos uno de los
    idiomas de la lista si cargue -- solo falla si NINGUNO carga (ver
    test_tesseract_ocr_engine.py, que ejercita ese caso real con
    lang='spa' a solas). Por eso este texto de prueba esta en ingles:
    es lo que se puede verificar con reconocimiento real en este
    entorno; en la maquina de despliegue con 'tesseract-ocr-spa'
    instalado, un documento en espanol se reconoceria igual de bien."""
    raiz, origenes_dir = raiz_temporal
    ctx = bootstrap(raiz=str(raiz))
    proyecto = ctx.projects.crear_proyecto(nombre="Extraccion OCR", perfil_id="perfil_generico")
    pctx = abrir_proyecto(ctx, proyecto)

    from PIL import Image, ImageDraw, ImageFont

    archivo = origenes_dir / "escaneo.png"
    fuente = ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf", 40)
    imagen = Image.new("RGB", (500, 100), color="white")
    ImageDraw.Draw(imagen).text((10, 10), "HELLO WORLD", font=fuente, fill="black")
    imagen.save(archivo)

    extraido = _registrar_clasificar_y_extraer(pctx, proyecto.id, str(archivo))
    assert extraido.estado.value == "extraido"

    markdown = Path(extraido.ruta_extraccion).read_text(encoding="utf-8")
    assert "HELLO WORLD" in markdown


def test_extraer_imagen_en_blanco_por_ocr_deja_advertencia(raiz_temporal) -> None:
    raiz, origenes_dir = raiz_temporal
    ctx = bootstrap(raiz=str(raiz))
    proyecto = ctx.projects.crear_proyecto(nombre="Extraccion OCR en blanco", perfil_id="perfil_generico")
    pctx = abrir_proyecto(ctx, proyecto)

    from PIL import Image

    archivo = origenes_dir / "en_blanco.png"
    Image.new("RGB", (100, 40), color="white").save(archivo)

    extraido = _registrar_clasificar_y_extraer(pctx, proyecto.id, str(archivo))
    assert extraido.estado.value == "extraido"
    assert extraido.advertencias_extraccion


def test_extraer_doc_legado_via_conversion_libreoffice(raiz_temporal) -> None:
    """DOC/XLS legado: conversion previa con LibreOffice headless antes
    de reusar los extractores de DOCX/XLSX del Modulo 5 (README,
    'Proximo modulo' -- ya resuelto). Se genera el .doc de prueba con el
    propio LibreOffice (round-trip DOCX->DOC), igual que
    test_legacy_extractors.py."""
    import subprocess
    import tempfile

    import docx

    raiz, origenes_dir = raiz_temporal
    ctx = bootstrap(raiz=str(raiz))
    proyecto = ctx.projects.crear_proyecto(nombre="Extraccion DOC legado", perfil_id="perfil_generico")
    pctx = abrir_proyecto(ctx, proyecto)

    d = docx.Document()
    d.add_heading("Politica legada", level=1)
    d.add_paragraph("Contenido de un documento .doc clasico.")
    docx_path = origenes_dir / "politica.docx"
    d.save(str(docx_path))

    perfil = tempfile.mkdtemp(prefix="rag_content_studio_smoke_lo_profile_")
    try:
        subprocess.run(
            [
                "soffice",
                "--headless",
                "--norestore",
                f"-env:UserInstallation=file://{perfil}",
                "--convert-to",
                "doc",
                "--outdir",
                str(origenes_dir),
                str(docx_path),
            ],
            capture_output=True,
            text=True,
            timeout=120,
            check=True,
        )
    finally:
        shutil.rmtree(perfil, ignore_errors=True)
    archivo = origenes_dir / "politica.doc"

    extraido = _registrar_clasificar_y_extraer(pctx, proyecto.id, str(archivo))
    assert extraido.estado.value == "extraido"

    markdown = Path(extraido.ruta_extraccion).read_text(encoding="utf-8")
    assert "Politica legada" in markdown
    assert "Contenido de un documento .doc clasico." in markdown


def test_extraer_xls_legado_via_conversion_libreoffice(raiz_temporal) -> None:
    import subprocess
    import tempfile

    import openpyxl

    raiz, origenes_dir = raiz_temporal
    ctx = bootstrap(raiz=str(raiz))
    proyecto = ctx.projects.crear_proyecto(nombre="Extraccion XLS legado", perfil_id="perfil_generico")
    pctx = abrir_proyecto(ctx, proyecto)

    libro = openpyxl.Workbook()
    hoja = libro.active
    hoja.title = "Datos"
    hoja.append(["nombre", "area"])
    hoja.append(["Ana", "RH"])
    xlsx_path = origenes_dir / "datos.xlsx"
    libro.save(str(xlsx_path))

    perfil = tempfile.mkdtemp(prefix="rag_content_studio_smoke_lo_profile_")
    try:
        subprocess.run(
            [
                "soffice",
                "--headless",
                "--norestore",
                f"-env:UserInstallation=file://{perfil}",
                "--convert-to",
                "xls",
                "--outdir",
                str(origenes_dir),
                str(xlsx_path),
            ],
            capture_output=True,
            text=True,
            timeout=120,
            check=True,
        )
    finally:
        shutil.rmtree(perfil, ignore_errors=True)
    archivo = origenes_dir / "datos.xls"

    extraido = _registrar_clasificar_y_extraer(pctx, proyecto.id, str(archivo))
    markdown = Path(extraido.ruta_extraccion).read_text(encoding="utf-8")
    assert "| nombre | area |" in markdown
    assert "| Ana | RH |" in markdown


def test_extraer_audio_por_transcripcion_falla_por_falta_de_acceso_a_internet(raiz_temporal) -> None:
    """Modulo 7: la ruta TRANSCRIPCION ya no se rechaza con
    UnsupportedExtractionRouteError -- un MP3/MP4 se clasifica a
    TRANSCRIPCION (Modulo 4) y AudioVideoTranscriptionExtractor +
    FasterWhisperEngine lo procesan de verdad, cableados por
    app.bootstrap.abrir_proyecto().

    A diferencia de OCR (Modulo 6, donde 'eng' si esta instalado
    localmente), Whisper necesita descargar los pesos del modelo desde
    Hugging Face Hub la primera vez -- y el proxy de red de este
    entorno de desarrollo bloquea huggingface.co (confirmado con curl:
    403 Forbidden). No hay ningun idioma de respaldo que evite esto,
    a diferencia de Tesseract: aqui la unica salida honesta es
    TranscriptionEngineUnavailableError, nunca un crash. En una maquina
    con acceso a Hugging Face Hub (o con los pesos ya cacheados), el
    mismo documento se transcribiria con exito -- ver
    test_transcription_extractor.py para la logica de formato Markdown,
    ya probada con un motor falso."""
    raiz, origenes_dir = raiz_temporal
    ctx = bootstrap(raiz=str(raiz))
    proyecto = ctx.projects.crear_proyecto(nombre="Extraccion transcripcion", perfil_id="perfil_generico")
    pctx = abrir_proyecto(ctx, proyecto)

    archivo = origenes_dir / "entrevista.mp3"
    archivo.write_bytes(b"contenido binario simulado de audio")

    documento = pctx.documents.registrar_documento(proyecto_id=proyecto.id, ruta_origen=str(archivo))
    clasificado = pctx.documents.clasificar_documento(documento.id)
    assert clasificado.tipo_procesamiento.value == "transcripcion"

    from domain.documents.exceptions import TranscriptionEngineUnavailableError

    with pytest.raises(TranscriptionEngineUnavailableError):
        pctx.documents.extraer_documento(documento.id)
