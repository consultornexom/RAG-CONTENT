"""Prueba de humo del Modulo 9 (normalizacion). Usa una raiz temporal,
nunca la ruta real de datos -- mismo patron que
test_extraer_documento_extremo_a_extremo.py.

Modulo 8 (revision contra originales) todavia no existe (requiere GUI,
fuera de alcance hasta que el backend de Modulo 9/10 este completo). Por
eso aqui se avanza el documento de EXTRAIDO a REVISADO con el primitivo
de bajo nivel Documento.avanzar_a() directamente sobre el repositorio,
igual que ya hacen las pruebas unitarias de NormalizarDocumento -- no
porque sea el flujo real de produccion, sino porque Modulo 9 solo
necesita que el documento este REVISADO, sin necesitar el caso de uso
real (Modulo 8) que lo produzca.
"""
from __future__ import annotations

import shutil
from pathlib import Path

import docx
import pytest

from app.bootstrap import abrir_proyecto, bootstrap
from domain.documents.entities import EstadoDocumento
from infrastructure.database.document_repository import SqlAlchemyDocumentRepository

RAIZ_PRUEBA = Path("/tmp/rag_content_studio_smoke_test_normalizacion")


@pytest.fixture()
def raiz_temporal():
    if RAIZ_PRUEBA.exists():
        shutil.rmtree(RAIZ_PRUEBA)
    RAIZ_PRUEBA.mkdir(parents=True)

    # perfil_generico real, tal cual (sin reglas_normalizacion) --
    # nunca se edita el YAML publicado real del proyecto.
    perfiles_dir = RAIZ_PRUEBA / "perfiles" / "perfil_generico"
    perfiles_dir.mkdir(parents=True)
    shutil.copy(
        Path(__file__).resolve().parents[2] / "perfiles" / "perfil_generico" / "0.1.yaml",
        perfiles_dir / "0.1.yaml",
    )

    # Perfil de prueba propio de esta raiz temporal (nunca se toca el
    # repositorio real de perfiles) con reglas_normalizacion reales,
    # para probar que app.bootstrap las resuelve y aplica de verdad.
    perfil_prueba_dir = RAIZ_PRUEBA / "perfiles" / "perfil_prueba_normalizacion"
    perfil_prueba_dir.mkdir(parents=True)
    (perfil_prueba_dir / "0.1.yaml").write_text(
        """id: perfil_prueba_normalizacion
version: "0.1"
descripcion: "Perfil de prueba de humo con reglas de normalizacion reales."
reglas_normalizacion:
  - tipo_regla: eliminar_linea_repetida
    parametros:
      texto: "Confidencial"
  - tipo_regla: colapsar_saltos_de_linea
    parametros: {}
""",
        encoding="utf-8",
    )

    origenes_dir = RAIZ_PRUEBA / "_origenes_de_prueba"
    origenes_dir.mkdir()
    yield RAIZ_PRUEBA, origenes_dir
    shutil.rmtree(RAIZ_PRUEBA, ignore_errors=True)


def _avanzar_a_revisado(pctx, documento_id: str):
    """Sustituto de Modulo 8 (todavia sin GUI): usa el mismo repositorio
    real (SqlAlchemy, sobre la sesion ya abierta de proyecto.db) que usa
    ProjectContext, no un fake -- para que la prueba siga siendo de
    extremo a extremo."""
    repo = SqlAlchemyDocumentRepository(pctx.session)
    documento = repo.get(documento_id)
    assert documento is not None
    revisado = documento.avanzar_a(EstadoDocumento.EN_REVISION).avanzar_a(EstadoDocumento.REVISADO)
    repo.save(revisado)
    return revisado


def test_normalizar_documento_aplica_reglas_reales_del_perfil_y_guarda_en_06(raiz_temporal) -> None:
    raiz, origenes_dir = raiz_temporal
    ctx = bootstrap(raiz=str(raiz))
    proyecto = ctx.projects.crear_proyecto(
        nombre="Normalizacion con reglas", perfil_id="perfil_prueba_normalizacion"
    )
    pctx = abrir_proyecto(ctx, proyecto)

    archivo = origenes_dir / "politica.docx"
    d = docx.Document()
    d.add_heading("Politica de vacaciones", level=1)
    d.add_paragraph("Confidencial")
    d.add_paragraph("Contenido real de la politica.")
    d.add_paragraph("Confidencial")
    d.save(str(archivo))

    documento = pctx.documents.registrar_documento(proyecto_id=proyecto.id, ruta_origen=str(archivo))
    clasificado = pctx.documents.clasificar_documento(documento.id)
    extraido = pctx.documents.extraer_documento(clasificado.id)
    assert extraido.estado == EstadoDocumento.EXTRAIDO

    _avanzar_a_revisado(pctx, extraido.id)

    normalizado = pctx.documents.normalizar_documento(extraido.id)
    assert normalizado.estado == EstadoDocumento.NORMALIZADO
    assert normalizado.ruta_normalizado is not None
    assert Path(normalizado.ruta_normalizado).parent == Path(proyecto.ruta_carpeta) / "06_normalizados"
    assert normalizado.cambios_normalizacion  # la regla eliminar_linea_repetida si tuvo efecto

    markdown_normalizado = Path(normalizado.ruta_normalizado).read_text(encoding="utf-8")
    assert "Confidencial" not in markdown_normalizado
    assert "Contenido real de la politica." in markdown_normalizado

    # El original extraido (03_extracciones/) no se toca: sigue intacto
    # y con "Confidencial" presente, la normalizacion escribe aparte.
    markdown_extraido = Path(extraido.ruta_extraccion).read_text(encoding="utf-8")
    assert "Confidencial" in markdown_extraido

    # Recuperado de proyecto.db, no el objeto en memoria.
    recuperado = pctx.documents.obtener_documento(normalizado.id)
    assert recuperado is not None
    assert recuperado.estado == EstadoDocumento.NORMALIZADO
    assert recuperado.ruta_normalizado == normalizado.ruta_normalizado


def test_normalizar_documento_sin_reglas_en_el_perfil_no_falla_y_no_deja_cambios(raiz_temporal) -> None:
    """perfil_generico real (Anexo B) no declara reglas_normalizacion --
    app.bootstrap.abrir_proyecto() debe resolver una lista vacia sin
    fallar, y NormalizarDocumento debe devolver el Markdown intacto."""
    raiz, origenes_dir = raiz_temporal
    ctx = bootstrap(raiz=str(raiz))
    proyecto = ctx.projects.crear_proyecto(nombre="Normalizacion sin reglas", perfil_id="perfil_generico")
    pctx = abrir_proyecto(ctx, proyecto)

    archivo = origenes_dir / "datos.csv"
    archivo.write_text("nombre,area\nAna,RH\n", encoding="utf-8")

    documento = pctx.documents.registrar_documento(proyecto_id=proyecto.id, ruta_origen=str(archivo))
    clasificado = pctx.documents.clasificar_documento(documento.id)
    extraido = pctx.documents.extraer_documento(clasificado.id)
    _avanzar_a_revisado(pctx, extraido.id)

    normalizado = pctx.documents.normalizar_documento(extraido.id)
    assert normalizado.estado == EstadoDocumento.NORMALIZADO
    assert normalizado.cambios_normalizacion == []

    markdown_revisado = Path(extraido.ruta_extraccion).read_text(encoding="utf-8")
    markdown_normalizado = Path(normalizado.ruta_normalizado).read_text(encoding="utf-8")
    assert markdown_normalizado == markdown_revisado
